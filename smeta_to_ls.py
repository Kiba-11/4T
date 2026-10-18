# -*- coding: utf-8 -*-
"""
УПСС · Обработка сметных листов.  © KAS-11. Все права принадлежат KAS-11.

Преобразование локальных сметных расчётов (выгрузка Smeta.RU в Excel, форма 4б)
в плоский реестр позиций — лист «-ЛС-» (как в «Разработка УПСС → 4Т → обработка сметных листов»).

Запуск:
    python smeta_to_ls.py                      — окно (выбор файлов/папки)
    python smeta_to_ls.py <файлы или папки> -o результат.xlsx [--code 5-5-1-1] [--object "Наименование"]
    python smeta_to_ls.py <файлы или папки> -a готовый.xlsx   — дописать сметы в конец листа «-ЛС-»
    python smeta_to_ls.py --svod готовый.xlsx                 — собрать лист «Свод смет» из «-ЛС-»
    python smeta_to_ls.py --itog готовый.xlsx                 — собрать лист «Итог» (по Разделу/Подразделу КР)
    python smeta_to_ls.py --grunt готовый.xlsx                — грунт: колонки T…AC на «-ЛС-» и лист «Грунт»

Требуется: Python 3.8+, openpyxl (pip install openpyxl); для PDF — pdfplumber (pip install pdfplumber)
"""
import argparse
import glob
import hashlib
import hmac
import os
import re
import sys
import traceback
import warnings

warnings.filterwarnings("ignore", module="openpyxl")  # «Cannot parse header or footer» — колонтитулы Smeta.RU, на данные не влияет

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- авторство и пароль
COPYRIGHT = "© KAS-11. Все права принадлежат KAS-11."
APP_AUTHOR = "KAS-11"
# Хеш пароля на запуск (PBKDF2-SHA256, «соль$хеш»). Пусто — пароль задаётся при первом запуске.
PASSWORD_HASH = "d45e9c2b877dcf54$85679e11f62dd5b3ae6b1e597b7b71dc66387f27195f7522c6dfc035d98ef8d0"


def _pw_hash(pw, salt=None):
    salt = salt or os.urandom(8).hex()
    h = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), salt.encode(), 120000).hex()
    return f"{salt}${h}"


def check_password(pw):
    if not PASSWORD_HASH or "$" not in PASSWORD_HASH:
        return False
    salt = PASSWORD_HASH.split("$", 1)[0]
    return hmac.compare_digest(_pw_hash(pw, salt), PASSWORD_HASH)


def store_password(pw):
    """Записать новый пароль (его хеш) прямо в этот файл программы."""
    global PASSWORD_HASH
    new = _pw_hash(pw)
    path = os.path.abspath(__file__)
    src = open(path, encoding="utf-8").read()
    src2, n = re.subn(r'^PASSWORD_HASH = ".*"$', f'PASSWORD_HASH = "{new}"', src, count=1, flags=re.M)
    if n != 1:
        raise RuntimeError("не найдена строка PASSWORD_HASH в файле программы")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(src2)
    PASSWORD_HASH = new


def _stamp(wb):
    """Авторство в свойствах создаваемой книги Excel."""
    try:
        wb.properties.creator = APP_AUTHOR
        wb.properties.lastModifiedBy = APP_AUTHOR
        wb.properties.description = "Сформировано программой «УПСС · Обработка сметных листов». " + COPYRIGHT
    except Exception:  # noqa
        pass


# --------------------------------------------------------------------------- настройки
NDS_EQUIP = 22            # колонка O «НДС» для строк оборудования
EQUIP_KIND = "ИО"         # колонка N «Доп. сведения по оборудованию» (ИО / ТО)
SORT_BY_LS_NUMBER = True  # сортировать сметы по номеру ЛС

HEADERS = {
    "A": "Раздел КР", "B": "Подраздел КР", "C": "Блок", "D": "Раздел КР", "E": "Подраздел КР",
    "F": "№ ЛС", "G": "№ п/п", "H": "Шифр расценки и коды ресурсов", "I": "Наименование работ и затрат",
    "J": "Ед. изм.", "K": "Кол-во единиц", "L": "Всего затрат в базисном уровне цен, руб.",
    "M": "ВСЕГО затрат в текущем уровне цен, руб.", "N": "Доп. сведения по оборудованию", "O": "НДС",
    "P": "Группа монт. (авто)", "Q": "Итоги по ЛС", "R": "Оборудование",
}
GRUNT = {"T": "Перевозка грунта", "U": "Разработка в отвал", "V": "Разработка без погрузки",
         "W": "Разработка вручную", "X": "Разработка с погрузкой", "Y": "Засыпка",
         "Z": "МАТЕРИАЛ для засыпки", "AA": "Утилизация мусора", "AB": "Перевозка мусора",
         "AC": "Утилизация грунта"}
WIDTHS = {"A": 14.3, "B": 14.3, "C": 10, "D": 8, "E": 8, "F": 10, "G": 7.1, "H": 14.3, "I": 60,
          "J": 14.3, "K": 10, "L": 18, "M": 18, "N": 8, "O": 6, "P": 5.9, "Q": 19, "R": 8,
          "S": 4, "T": 10, "U": 10, "V": 10, "W": 10, "X": 10, "Y": 10, "Z": 13, "AA": 10,
          "AB": 10, "AC": 10}


# --------------------------------------------------------------------------- утилиты
def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("_x000D_", "").replace("\r", "").strip()


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).split("\n")[0].replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def ls_sort_key(ls_no):
    """02-01-03 < 02-01-03к2 < 02-01-04;  09-01-06 < 09-01-06/1 < 09-01-07."""
    m = re.match(r"\s*(\d+(?:-\d+)*)(?:\s*[/._]\s*(\d+))?(.*)$", ls_no)
    if not m:
        return ((10 ** 9,), 0, ls_no)
    base = tuple(int(x) for x in m.group(1).split("-"))
    return (base, int(m.group(2)) if m.group(2) else 0, m.group(3).strip().lower())


class Position:
    __slots__ = ("num", "code", "name", "unit", "qty", "base", "cur", "kind", "parent", "closed", "group", "smont", "eqinfo", "nds")

    def __init__(self, num, code, name, unit, qty, base=0.0, cur=0.0, kind="", parent=None):
        self.num, self.code, self.name, self.unit, self.qty = num, code, name, unit, qty
        self.base, self.cur, self.kind, self.parent = base, cur, kind, parent
        self.closed = False
        self.group = "3"
        self.smont = None      # по листу Source: True — монтажные, False — строительные
        self.eqinfo = None     # «ИО»/«ТО» из сметы
        self.nds = None        # ставка НДС из сметы


class Estimate:
    def __init__(self, path, sheet=""):
        self.path = path
        self.sheet = sheet
        self.ls_no = ""
        self.ls_name = ""
        self.object_name = ""
        self.excl = False        # смета на исключение (лист «…_искл»)
        self.items = []          # ("section"|"sub", name) | ("pos", Position)
        self.total_base = None   # «Всего по смете» из файла — для сверки
        self.total_cur = None
        self.hdr_equip = None    # шапка: «Оборудование», руб. (базис, текущ.)
        self.hdr_mont = None     # шапка: «Работы по монтажу оборудования»
        self.total_from_hdr = False  # итог взят из шапки (тыс. руб.) — допуск на округление
        self.body_total = None       # «Всего по смете» из тела
        self.body_equip = None       # строка «Оборудование» в итогах тела
        self.returned = 0
        self.report_no = ""          # номер общего отчёта (обычно шифр УПСС)
        self.has_body_title = False
        self.pos_tot = [0.0, 0.0]    # сумма напечатанных «Всего по позиции» (для поиска ошибок в самой смете)
        self.no_total = []           # расценки с расшифровкой, у которых не нашлась строка «Всего по позиции»
        self.warnings = []
        self.ls_full = ""            # № ЛС как в заголовке, без обрезки («02-01-06 искл»)

    @property
    def key(self):
        """№ ЛС в «-ЛС-» — полностью, как в заголовке сметы («02-01-06 искл», «02-01-03 доб.»)."""
        k = getattr(self, "ls_full", "") or self.ls_no
        if self.excl and not re.search(r"искл", k, re.I):
            k += " искл"                     # «искл» только в имени листа — добавляем к номеру
        return k

    @property
    def src(self):
        b = os.path.basename(self.path)
        return f"{b} / {self.sheet}" if self.sheet else b


# --------------------------------------------------------------------------- разбор
SKIP_SHEETS = {"Source", "SourceObSm", "SmtRes", "EtalonRes"}
RE_TITLE = re.compile(r"ЛОКАЛЬН\w*\s+(СМЕТН\w*\s+РАСЧ[ЕЁ]Т|СМЕТА)\s*№\s*(.*)$", re.I)
RE_SECTION = re.compile(r"^Раздел\s*:?\s*(?:\d+[.:]?\s*)?(.*)$", re.I)
RE_SECTION_KR = re.compile(r"^Раздел\s*\d+(?:\.\d+)*\.", re.I)   # «Раздел 16.» — только такие идут в A/B


class SecName(str):
    """Название раздела + признак «Раздел N.» (для колонок A/B)."""
    kr = False


# ── Справочник «Раздел КР|Подраздел КР»: код ────────────────────────────────
# Колонки A/B заполняются ТОЛЬКО если раздел сметы («Раздел N. …») найден здесь.
# «ИТОГО» → подраздел пустой. Чтобы добавить вид — допишите строку в том же формате.
KR_MAP = {
    "Подготовка территории строительства|вырубка деревьев и кустарников": "1.4",
    "Подготовка территории строительства|компенсационное озеленение": "1.5",

    "Земляные работы|ИТОГО": "2.1.1",
    "Земляные работы|обратная засыпка": "2.1.1.1",
    "Крепление котлована|ИТОГО": "2.1.2",
    "Эксплуатация подъемно-транспортных механизмов|ИТОГО": "2.1.3",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|ИТОГО": "2.1.4",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|сваи": "2.1.4.1",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|фундамент": "2.1.4.2",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|колонны (пилоны)": "2.1.4.3",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|стены": "2.1.4.4",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|перекрытие, покрытие": "2.1.4.5",
    "Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий|лестницы": "2.1.4.6",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|ИТОГО": "2.1.5",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|колонны, пилоны: железобетонные": "2.1.5.1",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|колонны: металлические": "2.1.5.2",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|стены": "2.1.5.3",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|перекрытие, покрытие (крыша): железобетонные": "2.1.5.4",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|перекрытие, покрытие (крыша): металлические конструкции (фермы, балки и т.п.)": "2.1.5.5",
    "Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций|лестницы": "2.1.5.6",

    "Ненесущие стены и перегородки|ИТОГО": "2.1.6",
    "Ненесущие стены и перегородки|ненесущие стены": "2.1.6.1",
    "Ненесущие стены и перегородки|перегородки": "2.1.6.2",
    "Кровля|ИТОГО": "2.1.7",
    "Проемы:|оконные блоки, витражи, остекление лоджий и балконов": "2.1.8",
    "Проемы:|дверные блоки, ворота": "2.1.9",
    "Внутренняя отделка|ИТОГО": "2.1.10",
    "Внутренняя отделка|полы": "2.1.10.2",
    "Внутренняя отделка|потолки": "2.1.10.3",
    "Внутренняя отделка|стены": "2.1.10.1",
    "Наружная отделка|ИТОГО": "2.1.11",
    "Прочие конструктивные решения|ИТОГО": "2.1.12",

    "Электроснабжение, электроосвещение, молниезащита и заземление|ИТОГО": "2.2.1",
    "Система водоснабжения и водоотведения|ИТОГО": "2.2.2",
    "Система водоснабжения и водоотведения|": "2.2.2.1",
    "Система отопления|ИТОГО": "2.2.3",
    "Система отопления|": "2.2.3.1",
    "Система отопления|индивидуальный тепловой пункт (ИТП)": "2.2.3.2",
    "Системы вентиляции, кондиционирования|ИТОГО": "2.2.4",
    "Сети связи (внутренние)|ИТОГО": "2.2.5",
    "Системы безопасности|ИТОГО": "2.2.6",
    "Автоматизированная система управления инженерными системами|ИТОГО": "2.2.7",
    "Автоматизированная система коммерческого учета энергоснабжающих ресурсов|ИТОГО": "2.2.8",
    "Система газоснабжения|ИТОГО": "2.2.9",
    "Система газоанализации|ИТОГО": "2.2.10",
    "Система холодоснабжения|ИТОГО": "2.2.11",
    "Лифтовое оборудование|ИТОГО": "2.2.12",

    "Технологическое оборудование|ИТОГО": "2.3.1.1",
    "ТО 10|": "2.3.1.4",
    "ТО 22|": "2.3.1.2",
    "ТО 0|": "2.3.1.3",

    "Объекты подсобного и обслуживающего назначения|ИТОГО": "3.1",
    "Объекты энергетического хозяйства|ИТОГО": "4.1",
    "Объекты транспортного хозяйства и связи|ИТОГО": "5.1",
    "Наружные сети и сооружения водоснабжения, водоотведения, теплоснабжения и газоснабжения|ИТОГО": "6.1",
    "Благоустройство и озеленение территории|ИТОГО": "7.1",
    "посадка деревьев, кустарников|ИТОГО": "7.5",
    "Подпорные стены (гл. 7)|ИТОГО": "7.7",
    "Временные здания и сооружения|Конструктивные решения (основание кранов)": "8.2",
    "Временные здания и сооружения|Организация дорожного движения на период строительства": "8.3",
    "Основание под башенный кран|ИТОГО": "8.2",
    "ОДД на период строительства|ИТОГО": "8.3",
    "Прочие работы и затраты|ИТОГО": "9",
    "Прочие работы и затраты|затраты на перевозку отходов строительства и сноса, в том числе грунта автотранспортными средствами; затраты на прием отходов строительства и сноса отходополучателями": "9.2",
    "Прочие работы и затраты|затраты на проведение пусконаладочных работ": "9.1",
    # из шаблона листа «Итог» (без кода)
    "Подготовка территории строительства|ИТОГО": "",
    "Проемы:|ИТОГО": "",
    "Водоотлив (дренаж)|ИТОГО": "",
    "Временные здания и сооружения|ИТОГО": "",
    "Мусоропровод|ИТОГО": "",
    "Подпорные стены|ИТОГО": "",
    "Пусконаладочные работы|ИТОГО": "",
    "Система взимания платы|ИТОГО": "",
    "Цементация грунтов|ИТОГО": "",
}


def _kr_norm(s):
    s = str(s).lower().replace("ё", "е")
    s = re.sub(r"\s*([:;,.()])\s*", r"\1", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip(" .:;-–—")


_KR_LIST = []
for _k in KR_MAP:
    _a, _, _b = _k.partition("|")
    _KR_LIST.append((_a.strip(), "" if _b.strip().upper() == "ИТОГО" else _b.strip(),
                     _kr_norm(_a), _kr_norm(_b) if _b.strip().upper() != "ИТОГО" else ""))


def kr_lookup(name):
    """Раздел сметы → (Раздел КР, Подраздел КР) по KR_MAP; не найден → ("", "")."""
    n = _kr_norm(name)
    best_pair, best_a = None, None
    for a, b, na, nb in _KR_LIST:
        if not na or not n.startswith(na):
            continue
        rest = n[len(na):].lstrip(" :.,;-–—")
        if nb:
            if rest.startswith(nb) and (best_pair is None or len(na) + len(nb) > best_pair[0]):
                best_pair = (len(na) + len(nb), a, b)
        elif best_a is None or len(na) > best_a[0]:
            best_a = (len(na), a)
    if best_pair:
        return best_pair[1], best_pair[2]
    if best_a:
        return best_a[1], ""
    return "", ""
RE_POSNUM = re.compile(r"^\d+(?:[,./]\d+)?$")
NOT_SUB = re.compile(r"^(итог|всего|в том числе|локальная смета|объем|объём|составил|проверил|"
                     r"накладные|сметная прибыль|кроме того|справочно|\()", re.I)


def read_source_kinds(wb):
    """Вид затрат по номеру позиции из служебного листа Source (Smeta.RU)."""
    kinds = {}
    if "Source" not in wb.sheetnames:
        return kinds
    s = wb["Source"]
    for row in s.iter_rows(min_row=1, max_col=140, values_only=True):
        if row[0] in (17, 18) and len(row) > 136:
            n = txt(row[4])
            eg = txt(row[136])   # колонка EG — вид затрат
            bi = row[60]         # колонка BI — 1 строительные, 2 монтажные
            if n and n not in kinds and (eg or bi):
                kinds[n] = (eg, bi)
    return kinds


class Grid:
    """Лист, прочитанный целиком в память (быстро даже для десятков тысяч строк)."""

    def __init__(self, ws):
        self.title = ws.title
        try:
            ws.reset_dimensions()        # некоторые программы пишут неверный размер листа
        except AttributeError:
            pass
        self.rows = [list(r) for r in ws.iter_rows(values_only=True)]
        self.max_row = len(self.rows)
        self.max_column = max((len(r) for r in self.rows), default=0)

    def v(self, r, c):
        if 1 <= r <= self.max_row:
            row = self.rows[r - 1]
            if 1 <= c <= len(row):
                return row[c - 1]
        return None


def find_header_row(g, limit=80):
    for r in range(1, min(g.max_row, limit) + 1):
        for x in g.rows[r - 1]:
            if isinstance(x, str) and "Шифр расценки" in x:
                return r
    return None


def map_columns(g, hdr_row):
    m = {}
    for c in range(1, g.max_column + 1):
        v = g.v(hdr_row, c)
        if not isinstance(v, str):
            continue
        t = " ".join(v.lower().replace("-\n", "").replace("- ", "").split())
        if t.startswith("№ п/п") and "num" not in m:
            m["num"] = c
        elif t.startswith("шифр"):
            m["code"] = c
        elif t.startswith("наименование"):
            m["name"] = c
        elif t.startswith("ед"):
            m["unit"] = c
        elif t.startswith("кол"):
            m["qty"] = c
        elif "текущ" in t and "всего" in t:
            m["cur"] = c
        elif t.startswith("всего") and ("2000" in t or "базис" in t):
            m["base"] = c
        elif t.startswith("всего затрат") and "base" in m and "cur" not in m:
            m["cur"] = c                    # Гранд-Смета, форма 1а: «Всего затрат, руб.» — текущий уровень
        elif t.startswith("доп. сведения по оборудованию"):
            m["eqinfo"] = c
    # «Ставка НДС» — в строке заголовка или в строке под ней
    for r in (hdr_row, hdr_row + 1):
        for c in range(1, g.max_column + 1):
            v = g.v(r, c)
            if isinstance(v, str) and "ставка ндс" in v.lower():
                m.setdefault("nds", c)
    need = {"num", "code", "name", "unit", "qty", "base", "cur"}
    if need - m.keys():
        raise ValueError("в заголовке не найдены колонки: " + ", ".join(sorted(need - m.keys())))
    return m


# Оборудование в формуле цены: 1,2% заготовительно-складских (*1,012) и/или транспорт 3% (*1,03);
# у материалов — 2% (*1,02) без транспорта. Пробелы убираем: в PDF формула переносится.
EQUIP_PRICE = re.compile(r"\*1[,.]012(?![\d\w])|\*1[,.]03(?![\d,.\w])|\*1[,.]042(?![\d\w])|индОб"
                         r"|1[,.]2%(?:Заг|ЗСР)|3%Трансп", re.I)     # 1,042 = 1,012 × 1,03; «индОб» — индекс на оборудование


def classify(code, name, kind_src):
    """-> 'equip' | 'mont' | 'work' | 'mat'"""
    k = (kind_src or "").lower()
    if k:
        if "оборудован" in k and "монтаж" not in k:
            return "equip"
        if "монтаж" in k and "материал" not in k:
            return "mont"
        if "материал" in k:
            return "mat"
        return "work" if re.match(r"^\d+\.\d+-", code) and not code.startswith("1.") else "mat"
    # без листа Source — по шифру и наименованию
    flat = re.sub(r"\s+", "", name + "|" + code).lower()
    if re.match(r"^ОБОРУДОВАНИЕ\s*:", name.upper()) or code.startswith("13.") \
            or (not re.match(r"^1\.\d+-\d+-\d+", code) and EQUIP_PRICE.search(re.sub(r"\s+", "", name))) \
            or "стоимостьоборудования" in flat or "переченьтехнологическогооборудования" in flat \
            or "тсн-2001.13" in flat:
        return "equip"
    if re.match(r"^4\.\d+-", code):
        return "mont"
    if re.match(r"^(?:[2-9]|1[0-24-9])\.\d+-", code):
        return "work"
    return "mat"


def _row_numbers(g, r, from_col):
    row = g.rows[r - 1] if 1 <= r <= g.max_row else []
    return [float(x) for x in row[from_col - 1:]
            if isinstance(x, (int, float)) and not isinstance(x, bool)]


RE_DOCNAME = re.compile(r"^Наименование работ документа\s*[-–—]\s*(.*)$", re.I | re.S)


def _clean_ls_no(s):
    s = re.sub(r"\(?\s*\bискл\b\.?\s*\)?", " ", s.strip(), flags=re.I)
    return " ".join(s.split())


def _finalize(est, by_source):
    """Пост-обработка одной сметы: пустые позиции, знак «искл», группы P, проверки."""
    est.ls_name = re.sub(r"^на\s+", "", est.ls_name.strip())                 # Гранд-Смета: «на 1-1-1.Устройство…»
    if est.ls_no and est.ls_name.startswith(est.ls_no):          # «01-01-01 Геодезические знаки»
        est.ls_name = est.ls_name[len(est.ls_no):].strip(" .-–")
    est.ls_name = est.ls_name.strip(" ,;")
    if est.body_total is not None:
        est.total_base, est.total_cur = est.body_total
        est.total_from_hdr = False
    if est.hdr_equip is None and est.body_equip is not None:
        est.hdr_equip = est.body_equip
    if est.returned:
        est.warnings.append(f"возврат материалов не включён (строк: {est.returned})")
    if est.no_total:
        est.warnings.append(f"не найдена строка «Всего по позиции» у {len(est.no_total)} поз. "
                            f"(№ {', '.join(est.no_total[:10])}{'…' if len(est.no_total) > 10 else ''}) — их суммы неполные")
    kept, zero = [], 0
    for t, o in est.items:
        if t == "pos" and o.qty == 0 and o.base == 0 and o.cur == 0:
            zero += 1
            continue
        kept.append((t, o))
    est.items = kept
    if zero:
        est.warnings.append(f"пропущено позиций с нулевым объёмом: {zero}")
    poss = [o for t, o in est.items if t == "pos"]
    if est.excl:
        if sum(p.cur for p in poss) > 0:
            for p in poss:
                p.base, p.cur = -p.base, -p.cur
            est.pos_tot = [-est.pos_tot[0], -est.pos_tot[1]]
            est.warnings.append("суммы сметы на исключение взяты со знаком «минус»")
        neg = lambda v: None if v is None else -abs(v)
        est.total_base, est.total_cur = neg(est.total_base), neg(est.total_cur)
        if est.hdr_equip:
            est.hdr_equip = tuple(-abs(v) for v in est.hdr_equip)
        if est.hdr_mont:
            est.hdr_mont = tuple(-abs(v) for v in est.hdr_mont)
    assign_groups(poss, by_source=by_source)
    if poss and all(p.base == 0 and p.cur == 0 for p in poss):
        est.warnings.append("все суммы нулевые — книга, вероятно, сохранена без пересчёта формул. "
                            "Откройте её в Excel и сохраните, затем повторите.")


def parse_sheet(g, path, kinds):
    """Разбор листа. Возвращает список смет: обычно одна, но в «отчёте» на одном листе
    может идти подряд много ЛС, каждая со своим заголовком «ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЕТ № …»."""
    hdr = find_header_row(g)
    if not hdr:
        return []
    # «Шифр расценки» может стоять в обрывке многострочной шапки — берём первую строку, где нашлись все графы
    cols, err = None, None
    for r in range(hdr, min(g.max_row, hdr + 40) + 1):
        if not any(isinstance(x, str) and "Шифр расценки" in x for x in g.rows[r - 1]):
            continue
        try:
            cols = map_columns(g, r)
            hdr = r
            break
        except ValueError as ex:
            err = err or ex
    if cols is None:
        raise err
    sheet_excl = "искл" in g.title.lower()
    est = Estimate(path, g.title)
    est.excl = sheet_excl

    # ---- шапка (текст может стоять в любой колонке)
    prev = ""
    for r in range(1, hdr):
        cells = [(c, txt(g.v(r, c))) for c in range(1, min(g.max_column, 14) + 1)]
        cells = [(c, t) for c, t in cells if t]
        if not cells:
            continue
        line = cells[0][1]
        low = line.lower()
        m = RE_TITLE.search(line)
        if m:
            est.ls_no = _clean_ls_no(m.group(2))
            est.ls_full = " ".join(m.group(2).split())
            if re.search(r"\bискл", m.group(2), re.I):
                est.excl = True
        elif low.startswith("(наименование стройки"):
            est.object_name = prev.strip(' "«»')
        elif low.startswith("(наименование работ"):
            est.ls_name = prev
        elif low.startswith("(локальная смета)"):
            pass
        elif low.startswith("объект:") or (len(cells) > 1 and low.startswith("объект")):
            est.object_name = cells[1][1] if len(cells) > 1 else line.split(":", 1)[-1].strip()
        else:
            label = " ".join(line.split()).rstrip(":")
            nums = _row_numbers(g, r, cells[0][0] + 1)
            if len(nums) >= 2 and label in ("Оборудование", "Работы по монтажу оборудования",
                                             "Сметная стоимость"):
                pair = (round(nums[0] * 1000, 2), round(nums[1] * 1000, 2))
                if label == "Оборудование":
                    est.hdr_equip = pair
                elif label.startswith("Работы по монтажу"):
                    est.hdr_mont = pair
                elif est.total_base is None:
                    est.total_base, est.total_cur = pair
                    est.total_from_hdr = True
            prev = line

    out = [est]
    c_num, c_code, c_name, c_unit, c_qty = cols["num"], cols["code"], cols["name"], cols["unit"], cols["qty"]
    c_base, c_cur = cols["base"], cols["cur"]
    c_eq, c_nds = cols.get("eqinfo"), cols.get("nds")
    open_pos, embedded = None, []
    in_breakdown = False      # внутри расшифровки расценки (ЗП/ЭМ/МР/НР/СП…)
    last_closed = None        # (позиция, её вложенные ресурсы, итог «Всего по позиции»)
    body_titles = 0

    for r in range(hdr + 1, g.max_row + 1):
        a = txt(g.v(r, c_num))
        code_raw = txt(g.v(r, c_code))
        name = txt(g.v(r, c_name))
        name_l = name.lower()
        a_l = a.lower()
        if a == "1" and code_raw == "2" and name == "3":     # строка с номерами граф
            continue
        # Гранд-Смета: пометки ресурсов вместо номера
        if a and code_raw and not RE_POSNUM.match(a):
            mk = a.replace(" ", "").lower()
            if mk in ("уд", "зам"):          # удалённый / заменённый ресурс расценки — уже учтён в ЗП/ЭМ/МР
                continue
            if open_pos is not None and re.fullmatch(r"(н|д|о|н,зам|зам,н|д,зам)", mk):
                a = f"{open_pos.num},{sum(1 for e in embedded) + 1}"   # неучтённый/добавленный ресурс

        # ---- новая ЛС внутри листа
        if a and not RE_POSNUM.match(a):
            m = RE_TITLE.search(a)
            if m:
                body_titles += 1
                no = _clean_ls_no(m.group(2))
                if est.items or body_titles > 1 and est.ls_no and est.has_body_title:
                    est = Estimate(path, g.title)
                    est.object_name = out[0].object_name
                    out.append(est)
                else:                         # первая ЛС: шапка листа описывала весь отчёт
                    if est.ls_no and est.ls_no != no:
                        est.report_no = est.ls_no
                    est.total_base = est.total_cur = None
                    est.hdr_equip = est.hdr_mont = None
                    est.total_from_hdr = False
                    est.ls_name = ""
                est.has_body_title = True
                est.ls_no = no
                est.ls_full = " ".join(m.group(2).split())
                est.excl = sheet_excl or bool(re.search(r"\bискл", m.group(2), re.I))
                open_pos, embedded, in_breakdown = None, [], False
                continue
            m = RE_DOCNAME.match(a)
            if m:
                est.ls_name = " ".join(m.group(1).split())
                continue

        if est.body_total is not None:
            continue          # после «Всего по смете» — приложения/итоги отчёта, не позиции этой ЛС
        if open_pos is not None and not a and re.match(r"^(ЗП|ЭМ|МР|ОТм?|в т\.\s*ч\.|НР|СП|ЗТР)", name):
            in_breakdown = True

        # итог по позиции
        # Smeta.RU иногда делит итог на «ВСЕГО работ по позиции» + «ВСЕГО оборудование по позиции»
        if re.match(r"^всего\s+оборудовани\w*\s+по\s+позиции", name_l) and last_closed is not None:
            vals = _row_numbers(g, r, c_name + 1)
            if len(vals) >= 2 and (vals[0] or vals[1]):
                lp, lemb, tot = last_closed
                tot = (tot[0] + vals[0], tot[1] + vals[1])
                est.pos_tot[0] += vals[0]
                est.pos_tot[1] += vals[1]
                lp.base = round(tot[0] - sum(e.base for e in lemb), 2)
                lp.cur = round(tot[1] - sum(e.cur for e in lemb), 2)
                last_closed = (lp, lemb, tot)
            continue
        if re.match(r"^всего\s+(?:работ\s+)?по\s+позиции", name_l) and open_pos is not None:
            vals = _row_numbers(g, r, c_name + 1)
            if len(vals) >= 2:
                open_pos.base = round(vals[0] - sum(e.base for e in embedded), 2)
                open_pos.cur = round(vals[1] - sum(e.cur for e in embedded), 2)
                last_closed = (open_pos, embedded, (vals[0], vals[1]))
                est.pos_tot[0] += vals[0]
                est.pos_tot[1] += vals[1]
            open_pos.closed = True
            open_pos, embedded, in_breakdown = None, [], False
            continue

        # итог по смете (берём последнюю строку с числами)
        if re.match(r"^(всего по смете|итого по локальной смете|итого по смете)", name_l) \
                or re.match(r"^(всего по смете|итого по локальной смете|итого по смете)", a_l):
            vals = _row_numbers(g, r, (c_name if name else c_num) + 1)
            if len(vals) >= 2:
                est.body_total = (vals[0], vals[1])
            continue
        # строка «Оборудование» в итогах сметы
        lab = (a_l or name_l).strip().rstrip(":")
        if lab == "оборудование" and not code_raw:
            if est.body_total is not None:        # после «Всего по смете» — общий итог отчёта, не этой ЛС
                continue
            vals = _row_numbers(g, r, (c_num if a else c_name) + 1)
            if len(vals) >= 2:
                est.body_equip = (vals[0], vals[1])
            continue

        # позиция / ресурс
        if a and (code_raw or name) and RE_POSNUM.match(a):
            if name.upper().startswith("ВОЗВРАТ"):            # возврат материалов — «кроме того», в итог не входит
                est.returned += 1
                open_pos = None
                continue
            lines = [x.strip() for x in code_raw.split("\n") if x.strip()]
            code = lines[0] if lines else code_raw
            full_name = name + (f"\n({'; '.join(lines[1:])})" if len(lines) > 1 else "")
            pos = Position(a, code, full_name, txt(g.v(r, c_unit)).split("\n")[0],
                           num(g.v(r, c_qty)), num(g.v(r, c_base)), num(g.v(r, c_cur)))
            eq = txt(g.v(r, c_eq)) if c_eq else ""
            if eq:                                            # «ИО» / «ТО» прямо в смете
                pos.kind = "equip"
                pos.eqinfo = eq
                pos.nds = g.v(r, c_nds) if c_nds else None
            else:
                src = kinds.get(a)
                pos.kind = classify(code, name, src[0] if src else None)
                if c_eq and pos.kind == "equip":     # в смете есть колонка ИО/ТО, а пометки нет — не оборудование
                    pos.kind = "mat"
                if src and src[1] in (1, 2):
                    pos.smont = src[1] == 2
            if open_pos is not None and ("," in a or in_breakdown):   # неучтённый ресурс внутри расценки
                pos.parent = open_pos
                embedded.append(pos)
            else:
                if open_pos is not None and in_breakdown and not open_pos.closed:
                    est.no_total.append(open_pos.num)
                open_pos, embedded, in_breakdown = pos, [], False
            est.items.append(("pos", pos))
            continue

        if a and not code_raw:
            if RE_SECTION.match(a) and a_l.startswith("раздел"):
                sec = SecName(RE_SECTION.match(a).group(1).strip() or a)
                sec.kr = bool(RE_SECTION_KR.match(a))     # «Раздел 16. …» → заполняем A/B
                est.items.append(("section", sec))
                open_pos = None
            elif a_l.startswith("подраздел"):
                est.items.append(("sub", re.sub(r"^Подраздел\s*:?\s*", "", a, flags=re.I).strip()))
            elif not NOT_SUB.match(a) and not RE_POSNUM.match(a) and not _row_numbers(g, r, c_num + 1):
                est.items.append(("sub", a))

    out = [e for e in out if e.items or e.body_total is not None or len(out) == 1]
    for e in out:
        if not e.ls_no:
            m = re.search(r"\d{2}-\d{2}-\d{2}(?:[.\-_]\d+)?", g.title) or \
                re.search(r"\d{2}-\d{2}(?:-\d+)?", os.path.basename(path))
            e.ls_no = m.group(0) if m else g.title
            e.warnings.append("номер ЛС не найден в шапке, взят из имени листа/файла")
        if not e.report_no and out[0].report_no:
            e.report_no = out[0].report_no
        _finalize(e, by_source=bool(kinds))
    return out


# --------------------------------------------------------------------------- PDF
class RowsGrid(Grid):
    """Таблица, собранная не из Excel, а из готовых строк (например, из PDF)."""

    def __init__(self, title, rows):  # noqa  (не вызываем Grid.__init__)
        self.title = title
        self.rows = rows
        self.max_row = len(rows)
        self.max_column = max((len(r) for r in rows), default=0)


RE_NUMCELL = re.compile(r"^[-−–]?\s?\d[\d\s\xa0]*(?:,\d+)?$")     # «-476,17», «- 476,17», «- 5 684,00»
PDF_HEADER = ["№ п/п", "Шифр расценки и коды ресурсов", "Наименование работ и затрат", "Ед. изм.",
              "Кол-во единиц", "Цена на ед. изм., руб.", "Поправочные коэффициенты",
              "Коэффициенты зимних удорожаний", "Всего затрат в базисном уровне цен, руб.",
              "Коэффициенты (индексы) пересчета", "ВСЕГО затрат в текущем уровне цен, руб."]


def _pdf_num(s):
    s = s.strip()
    if RE_NUMCELL.match(s):
        try:
            return float(re.sub(r"[\s\xa0]", "", s).replace("−", "-").replace("–", "-").replace(",", "."))
        except ValueError:
            pass
    return s


def _pdf_lines(words):
    """Слова → строки. Слово относится к строке, если перекрывается с ней по высоте больше чем наполовину
    (цифра, набранная другим шрифтом чуть ниже, остаётся в своей строке)."""
    lines = []   # [top, bottom, [words]]
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        h = max(w["bottom"] - w["top"], 0.1)
        placed = False
        for ln in reversed(lines[-3:]):
            ov = min(ln[1], w["bottom"]) - max(ln[0], w["top"])
            if ov >= 0.5 * min(h, ln[1] - ln[0]):
                ln[2].append(w)
                ln[0], ln[1] = min(ln[0], w["top"]), max(ln[1], w["bottom"])
                placed = True
                break
        if not placed:
            lines.append([w["top"], w["bottom"], [w]])
    lines.sort(key=lambda ln: ln[0])
    return [sorted(ln[2], key=lambda w: w["x0"]) for ln in lines]


def _wjoin(ws):
    """Склеить слова: вплотную стоящие куски — без пробела («01-0»+«4»+«-01» → «01-04-01»)."""
    out = ""
    prev = None
    for w in ws:
        if prev is not None and w["x0"] - prev["x1"] > 0.8:
            out += " "
        out += w["text"]
        prev = w
    return out


def _merge_num_words(line):
    """«3» «702,73» → «3 702,73» (пробел — разделитель тысяч); «-» «476,17» → «- 476,17»."""
    out = []
    for w in line:
        if out and out[-1]["text"] in ("-", "−", "–") and re.match(r"^\d", w["text"]) \
                and w["x0"] - out[-1]["x1"] < 6:
            m = dict(out[-1])
            m["text"] = "-" + w["text"]
            m["x1"] = w["x1"]
            out[-1] = m
        elif out and RE_NUMCELL.match(out[-1]["text"]) and re.fullmatch(r"\d{3}(?:,\d+)?", w["text"]) \
                and re.fullmatch(r"-?\d{1,3}(?: \d{3})*", out[-1]["text"]) and w["x0"] - out[-1]["x1"] < 5:
            m = dict(out[-1])
            m["text"] = out[-1]["text"] + " " + w["text"]
            m["x1"] = w["x1"]
            out[-1] = m
        else:
            out.append(w)
    return out


def _gap_cells(line, gap=12):
    """Строка шапки → ячейки по большим промежуткам; каждое число — отдельная ячейка."""
    line = _merge_num_words(line)
    cells, cur, last_x1 = [], [], None
    for w in line:
        isnum = bool(RE_NUMCELL.match(w["text"]))
        if cur and (w["x0"] - last_x1 > gap or isnum or RE_NUMCELL.match(cur[-1]["text"])):
            cells.append(_wjoin(cur))
            cur = []
        cur.append(w)
        last_x1 = w["x1"]
    if cur:
        cells.append(_wjoin(cur))
    return [_pdf_num(c) for c in cells]


def _col_bounds(page, colnum_top, centers):
    xs = sorted(e["x0"] for e in page.edges
                if e["orientation"] == "v" and e["top"] <= colnum_top + 3 and e["bottom"] >= colnum_top - 3)
    cl = []
    for x in xs:
        if not cl or x - cl[-1] > 3:
            cl.append(x)
    if len(cl) == len(centers) + 1:
        return cl
    # запасной вариант — середины между номерами граф
    b = [centers[0] - (centers[1] - centers[0]) / 2]
    b += [(centers[i] + centers[i + 1]) / 2 for i in range(len(centers) - 1)]
    b.append(centers[-1] + (centers[-1] - centers[-2]) / 2)
    return b


_SHORT_WORDS = {"по", "в", "во", "на", "и", "с", "со", "до", "от", "к", "для", "из", "без", "при", "за", "под"}


def _join(a, b, unit=False, code=False):
    if not a:
        return b
    if a.endswith("-") and not a.endswith(" -"):
        return a + b
    tail = a.split(" ")[-1]
    if (unit or code) and tail[-1:].isalpha() and tail[-1:].islower() and b[:1].islower() \
            and ((code and len(tail) >= 4) or (unit and len(tail) <= 7)) \
            and b.split(" ")[0] not in _SHORT_WORDS:
        return a + b                       # перенос слова без дефиса в узкой графе
    return a + " " + b


PDF_SKIP = re.compile(r"^(©|Программный комплекс|Smeta\.RU|Гранд-?Смета|\d{1,4}$)", re.I)


def pdf_to_grids(path, log=None):
    """PDF со сметами (форма 4 / 4б: «Строительный эксперт», Smeta.RU и т.п.) → список таблиц, по одной на ЛС."""
    try:
        import pdfplumber
    except ImportError:
        raise ValueError("для PDF нужен пакет pdfplumber: pip install pdfplumber")
    grids, cur = [], None
    bounds = None
    prev_kind = None      # 'pos' — позиция или её продолжение; 'section' — заголовок раздела
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        for pi, page in enumerate(pdf.pages):
            if log and pi and pi % 100 == 0:
                log(f"   … PDF: прочитано страниц {pi} из {n}")
            words = page.extract_words(x_tolerance=1.5, y_tolerance=2)
            lines = _pdf_lines(words)
            texts = [_wjoin(l) for l in lines]
            # строка с номерами граф 1…11
            colnum_idx = None
            for i, l in enumerate(lines):
                if [w["text"] for w in l] == [str(k) for k in range(1, 12)]:
                    colnum_idx = i
                    centers = [(w["x0"] + w["x1"]) / 2 for w in l]
                    bounds = _col_bounds(page, l[0]["top"], centers)
                    break
            title_idx = next((i for i, t in enumerate(texts[:colnum_idx] if colnum_idx is not None else texts)
                              if RE_TITLE.search(t)), None)
            body_from = 0
            if title_idx is not None and colnum_idx is not None:
                # новая ЛС: всё до шапки таблицы — титул
                hdr_lines = [(t, l) for t, l in zip(texts[:colnum_idx], lines[:colnum_idx]) if not PDF_SKIP.match(t)]
                rows = _pdf_title_rows(hdr_lines)
                rows.append(list(PDF_HEADER))
                cur = RowsGrid(f"PDF стр. {pi + 1}", rows)
                grids.append(cur)
                body_from = colnum_idx + 1
                prev_kind = None
            elif colnum_idx is not None:
                body_from = colnum_idx + 1
            if cur is None or bounds is None:
                continue
            for l, t in zip(lines[body_from:], texts[body_from:]):
                if PDF_SKIP.match(t) or re.match(r"^Объ[её]м\s*:", t):
                    continue
                l = _merge_num_words(l)
                # заголовки разделов (с первой графы или по центру)
                if re.match(r"^(Раздел|Подраздел)\b", t) and not re.search(r"\d+,\d{2}\s*$", t):
                    cur.rows.append([t])
                    prev_kind = "section"
                    continue
                if re.match(r"^Локальная смета\s*:", t, re.I):
                    cur.rows.append([t])
                    prev_kind = None
                    continue
                cells = [""] * 11
                for w in l:
                    x = w["x0"] + 1
                    c = 0
                    for k in range(11):
                        if x >= bounds[k]:
                            c = k
                    cells[c] = (cells[c] + " " + w["text"]).strip()
                has_num = any(RE_NUMCELL.match(cells[k]) for k in range(4, 11) if cells[k])
                # итоговые подписи могут начинаться левее графы 3 (в графах 1–2) — собираем подпись в графу 3
                if re.match(r"^(Всего по позиции|Всего по смете|Итого по локальной смете|Итого|Всего)\b", t, re.I) \
                        and not RE_POSNUM.match(cells[0] or "x"):
                    label = " ".join(cells[k] for k in range(0, 4) if cells[k] and not RE_NUMCELL.match(cells[k]))
                    cells[0] = cells[1] = cells[3] = ""
                    cells[2] = label
                # продолжение названия раздела на следующей строке
                if prev_kind == "section" and not has_num and not RE_POSNUM.match(cells[0] or "x") \
                        and not re.match(r"^(Раздел|Подраздел|Итого|Всего|Локальная)", t):
                    cur.rows[-1][0] = _join(cur.rows[-1][0], t)
                    continue
                last = cur.rows[-1] if cur.rows else None
                # продолжение наименования / ед. изм. / шифра на следующей строке
                if not cells[0] and not has_num and any(cells[k] for k in (1, 2, 3)) \
                        and not any(cells[k] for k in range(4, 11)) and prev_kind == "pos":
                    for k in (1, 2, 3):
                        if cells[k]:
                            last[k] = _join(str(last[k] or ""), cells[k], unit=(k == 3), code=(k == 1))
                    continue
                # Smeta.RU: ед. изм., количество и цены — на последней строке многострочной позиции
                if not cells[0] and has_num and prev_kind == "pos" and last is not None \
                        and not any(last[k] not in ("", None) for k in range(4, 11)) \
                        and not re.match(r"^(ЗП|ЭМ|МР|ОТ|в т\.|НР|СП|ЗТР|Всего|Итого)", cells[2]):
                    for k in range(11):
                        if not cells[k]:
                            continue
                        if k <= 3:
                            last[k] = _join(str(last[k] or ""), cells[k], code=(k == 1), unit=(k == 3))
                        else:
                            last[k] = _pdf_num(cells[k]) if k >= 4 else cells[k]
                    prev_kind = "pos_done"
                    continue
                row = [_pdf_num(c) if k >= 4 else c for k, c in enumerate(cells)]
                label = cells[2].lower().rstrip(":")
                if not any(cells[k] for k in range(0, 8)) and isinstance(row[8], float) \
                        and isinstance(row[10], float) and not cells[9]:
                    row[2] = "Всего по позиции"          # жирная строка итога позиции без подписи
                elif label == "всего" and has_num:
                    row[2] = "Всего по смете"
                elif label.startswith("оборудования") and has_num:
                    row[2] = "Оборудование"
                cur.rows.append(row)
                if cells[0]:
                    prev_kind = "pos"
                elif has_num or cells[2]:
                    prev_kind = None
            cur.max_row = len(cur.rows)
            cur.max_column = max(cur.max_column, 11)
    return grids


def _pdf_title_rows(hdr_lines):
    """Титул ЛС из PDF: склеиваем многострочные наименования объекта и работ."""
    texts = [t for t, _ in hdr_lines]
    rows = []
    i = 0
    while i < len(hdr_lines):
        t, l = hdr_lines[i]
        low = t.lower()
        if low.startswith("форма"):
            # наименование стройки — всё до «(наименование стройки…)»
            j = i + 1
            buf = []
            while j < len(texts) and not texts[j].lower().startswith("(наименование стройки"):
                buf.append(texts[j])
                j += 1
            if buf:
                rows.append([" ".join(buf)])
            i = j
            continue
        if RE_TITLE.search(t) or low.startswith("(локальная смета)"):
            rows.append([t])
            j = i + 1
            if j < len(texts) and texts[j].lower().startswith("(локальная смета)"):
                j += 1
            buf = []
            while j < len(texts) and not texts[j].lower().startswith("(наименование работ"):
                buf.append(texts[j])
                j += 1
            if buf:
                rows.append([" ".join(buf)])
            i = j
            continue
        rows.append(_gap_cells(l))
        i += 1
    return rows


def parse_file(path, log=None):
    """Все сметы файла: по листам, а внутри листа — по заголовкам ЛС. Порядок — как в книге."""
    if path.lower().endswith(".pdf"):
        out, errors = [], []
        for g in pdf_to_grids(path, log):
            try:
                out += parse_sheet(g, path, {})
            except Exception as ex:  # noqa
                errors.append(f"{g.title}: {ex}")
        return out, errors
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    grids = []
    for ws in wb.worksheets:
        if ws.title in SKIP_SHEETS:
            continue
        g = Grid(ws)
        if find_header_row(g):
            grids.append(g)
    # лист Source описывает одну смету — используем его, только если лист сметы один
    kinds = read_source_kinds(wb) if len(grids) == 1 else {}
    wb.close()
    out, errors = [], []
    for g in grids:
        try:
            out += parse_sheet(g, path, kinds)
        except Exception as ex:  # noqa
            errors.append(f"{g.title}: {ex}")
    return out, errors


def parse_estimate(path):
    """Совместимость: первая смета файла."""
    ests, errors = parse_file(path)
    if not ests:
        raise ValueError("; ".join(errors) or "не найден лист сметы со строкой заголовка «Шифр расценки…»")
    return ests[0]


# --------------------------------------------------------------------------- вывод
def assign_groups(poss, by_source=False):
    """Колонка P «Группа монт. (авто)»: 4 — монтаж, иначе 3/5/6 по сборнику расценки.
    Материалы и неучтённые ресурсы наследуют группу предыдущей расценки (как в 4Т).
    Если есть лист Source (Smeta.RU) — материалы берут группу по виду затрат сметы."""
    last = "3"
    for p in poss:
        if by_source and p.smont is not None and p.kind != "equip":
            if p.smont:
                g = "4"
            else:
                m = re.match(r"^(\d+)\.", p.code)
                g = m.group(1) if p.kind == "work" and m and m.group(1) in ("5", "6") else "3"
        elif p.kind == "mont":
            g = "4"
        elif p.kind == "work":
            m = re.match(r"^(\d+)\.", p.code)
            g = m.group(1) if m and m.group(1) in ("3", "5", "6") else "3"
        elif p.kind == "mat" and by_source:
            g = "3"
        else:                       # материал / оборудование — как у предыдущей расценки
            g = last
        p.group = g
        if p.parent is None and p.kind in ("mont", "work"):
            last = g


BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
LS_FILL = PatternFill("solid", fgColor="FFF2CC")
EXCL_FILL = PatternFill("solid", fgColor="FCE4D6")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
ITOG_MISMATCH_FILL = PatternFill("solid", fgColor="FFFFC7CE")
ITOG_MISMATCH_FONT = Font(color="FF9C0006", bold=True)
SEC_FILL = PatternFill("solid", fgColor="E2EFDA")
NEW_FILL = PatternFill("solid", fgColor="F8CBAD")
_THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
LS_SHEET = "-ЛС-"
PROTO_SHEET = "Протокол"


class _NullCell:
    """Заглушка для колонки, которой нет на листе: запись игнорируется."""
    value = font = fill = alignment = number_format = border = None


def _norm(t):
    return " ".join(str(t).lower().replace("ё", "е").replace("\n", " ").split())


class LSSheet:
    """Лист «-ЛС-», у которого колонки ищутся по заголовкам.
    Код пишет в «стандартные» колонки шаблона (F=6, G=7 … R=18, T…AC — грунт),
    обёртка переводит их в фактические колонки листа."""

    def __init__(self, ws, cmap=None, hdr=3):
        self.ws, self.hdr = ws, hdr
        self.cmap = cmap or {c: c for c in range(1, 30)}
        self.missing = []

    @classmethod
    def detect(cls, ws):
        try:
            hdr = find_ls_columns(ws)[0]
        except ValueError:
            return cls(ws)                       # лист пустой/новый — стандартная раскладка
        std = {openpyxl.utils.column_index_from_string(k): _norm(v) for k, v in HEADERS.items()}
        for c in range(1, ws.max_column + 1):          # старое название колонки C
            if _norm(ws.cell(hdr, c).value or "") == _norm("Примечание"):
                ws.cell(hdr, c).value = "Блок"
        fact = [(c, _norm(ws.cell(hdr, c).value)) for c in range(1, ws.max_column + 1)
                if ws.cell(hdr, c).value not in (None, "")]
        used, cmap, missing = set(), {}, []
        for c_std, title in sorted(std.items()):
            key = title[:25]
            hit = next((c for c, t in fact if c not in used and (t.startswith(key) or key.startswith(t[:25]))), None)
            if hit:
                cmap[c_std] = hit
                used.add(hit)
            else:
                cmap[c_std] = None
                missing.append(HEADERS[get_column_letter(c_std)])
        # колонки грунта: заголовки в шапке или в строках над ней
        for k, title in GRUNT.items():
            c_std = openpyxl.utils.column_index_from_string(k)
            hit = None
            for r in range(1, hdr + 1):
                for c in range(1, ws.max_column + 1):
                    if _norm(ws.cell(r, c).value or "") == _norm(title):
                        hit = c
                        break
                if hit:
                    break
            cmap[c_std] = hit
        obj = cls(ws, cmap, hdr)
        obj.missing = missing
        return obj

    def col(self, c):
        return self.cmap.get(c, c)

    def letter(self, c):
        x = self.col(c)
        return get_column_letter(x) if x else None

    def cell(self, r, c, value=None):
        x = self.col(c)
        if not x:
            return _NullCell()
        return self.ws.cell(r, x, value) if value is not None else self.ws.cell(r, x)

    @property
    def data_cols(self):
        return [x for x in self.cmap.values() if x]


def _init_ls_sheet(ws, upss_code, object_name):
    ws["A1"] = "Откорректировано"
    ws["G1"] = upss_code
    ws["N1"] = "Итоговая сумма"
    ws["G2"] = object_name
    ws["N2"] = "Итого по смете"
    for col, title in HEADERS.items():
        c = ws[f"{col}3"]
        c.value, c.font, c.fill, c.border = title, BOLD, HDR_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    _style_grunt_headers(ws, {openpyxl.utils.column_index_from_string(k): k for k in GRUNT}, 3)
    ws.row_dimensions[3].height = 60
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"


def _last_data_row(ls):
    """Последняя заполненная строка ниже шапки (по колонкам реестра)."""
    cols = [ls.col(c) for c in range(1, 19) if ls.col(c)]
    for r in range(ls.ws.max_row, ls.hdr, -1):
        for c in cols:
            if ls.ws.cell(r, c).value not in (None, ""):
                return r
    return ls.hdr


def existing_ls_numbers(ls):
    if not isinstance(ls, LSSheet):
        ls = LSSheet.detect(ls)
    nums = set()
    for r in range(ls.hdr + 1, ls.ws.max_row + 1):
        g = ls.cell(r, 7).value
        if isinstance(g, str) and RE_TITLE.search(g):
            nums.add(RE_TITLE.search(g).group(2).strip())
    return nums


def _split_section(name):
    """«Внутренняя отделка: полы» → («Внутренняя отделка», «полы»)."""
    depth = 0
    for i, ch in enumerate(name):
        if ch in "([«":
            depth += 1
        elif ch in ")]»":
            depth = max(0, depth - 1)
        elif ch == ":" and depth == 0 and not name[i + 1:i + 2] == "/":   # не делим по «https://»
            left, right = name[:i].strip(), name[i + 1:].strip()
            if left and right:
                return left, right
            break
    return name.strip(), ""


def _write_estimate(ws, r, est):
    F = est.key
    kr_a = kr_b = note = ""            # Раздел КР / Подраздел КР / Блок (подзаголовок)
    ls_row = r
    ws.cell(r, 6, F)
    ws.cell(r, 7, f"ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЕТ № {est.key}").font = BOLD
    for c in range(1, 19):
        ws.cell(r, c).fill = EXCL_FILL if est.excl else LS_FILL
    r += 1
    ws.cell(r, 6, F)
    ws.cell(r, 7, f"Наименование работ документа - {est.ls_name}"
                  + (" (смета на исключение)" if est.excl else ""))
    r += 1

    last_p = "3"
    sec_no, sec_row, sec_rows = 0, None, []
    for kind, obj in est.items:
        if kind == "section":
            if sec_row:
                sec_rows.append((sec_row, r - 1))
            sec_no += 1
            sec_row = r
            ws.cell(r, 6, F)
            ws.cell(r, 7, f"Раздел {sec_no}. {obj}").font = BOLD
            ws.cell(r, 16, last_p)
            # нет в справочнике (или не «Раздел N.») → раздел игнорируем, A/B остаются от предыдущего
            _a, _b = kr_lookup(obj) if getattr(obj, "kr", False) else ("", "")
            if _a:
                kr_a, kr_b = _a, _b
            note = ""
            for c in range(1, 19):
                ws.cell(r, c).fill = SEC_FILL
            r += 1
            continue
        if kind == "sub":
            note = obj
            ws.cell(r, 1, kr_a or None)
            ws.cell(r, 2, kr_b or None)
            ws.cell(r, 3, note)
            ws.cell(r, 6, F)
            ws.cell(r, 7, obj)
            ws.cell(r, 16, last_p)
            r += 1
            continue
        p = obj
        g = p.group
        if p.parent is None:
            last_p = g
        ws.cell(r, 1, kr_a or None)
        ws.cell(r, 2, kr_b or None)
        ws.cell(r, 3, note or None)
        ws.cell(r, 6, F)
        ws.cell(r, 7, p.num)
        ws.cell(r, 8, p.code)
        ws.cell(r, 9, p.name)
        ws.cell(r, 10, p.unit)
        ws.cell(r, 11, p.qty)
        ws.cell(r, 12, p.base).number_format = "#,##0.00"
        ws.cell(r, 13, p.cur).number_format = "#,##0.00"
        ws.cell(r, 16, g)
        if p.kind == "equip":
            ws.cell(r, 14, p.eqinfo or EQUIP_KIND)
            ws.cell(r, 15, p.nds if p.nds not in (None, "") else NDS_EQUIP)
            ws.cell(r, 18, "О")
        r += 1
    if sec_row:
        sec_rows.append((sec_row, r - 1))
    ws.cell(r, 6, F)
    ws.cell(r, 7, "В том числе (справочно)")
    ws.cell(r, 16, last_p)
    end = r
    M = ws.letter(13) if isinstance(ws, LSSheet) else "M"
    if M:
        ws.cell(ls_row, 17, f"=SUM({M}{ls_row + 1}:{M}{end})").number_format = "#,##0.00"
        for a, b in sec_rows:
            ws.cell(a, 17, f"=SUM({M}{a + 1}:{M}{b})").number_format = "#,##0.00"
    return r + 1


def _extend_total_formulas(ls, last):
    """Растянуть итоговые формулы строк 1–2 (L, M, T…AC) до новой последней строки."""
    ws = ls.ws
    rng = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)\s*:\s*(\$?[A-Z]{1,3}\$?)(\d+)")
    first = ls.hdr + 1
    for row in range(1, ls.hdr):
        for std in ["L", "M"] + list(GRUNT):
            col = ls.letter(openpyxl.utils.column_index_from_string(std))
            if not col:
                continue
            c = ws[f"{col}{row}"]
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                c.value = rng.sub(lambda m: f"{m.group(1)}{m.group(2)}:{m.group(3)}{max(int(m.group(4)), last)}", v)
            elif ls.hdr == 3 and v in (None, "") and (std in ("L", "M") or row == 2):
                func = "SUM" if (row == 2 and std in ("L", "M")) else "SUBTOTAL(9,"
                c.value = (f"=SUM({col}{first}:{col}{last})" if func == "SUM"
                           else f"=SUBTOTAL(9,{col}{first}:{col}{last})")
                c.number_format = "#,##0.00"
                c.font = BOLD
    # автофильтр
    ref = ws.auto_filter.ref
    if ref:
        m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", ref)
        if m:
            ws.auto_filter.ref = f"{m.group(1)}{m.group(2)}:{m.group(3)}{max(int(m.group(4)), last)}"
    else:
        ws.auto_filter.ref = f"A{ls.hdr}:{get_column_letter(max(ls.data_cols))}{last}"


def _write_protocol(wb, estimates, skipped, mark_new):
    head = ["Файл / лист", "№ ЛС", "Наименование ЛС", "Позиций",
            "Базис (реестр)", "Базис (смета)", "Текущ. (реестр)", "Текущ. (смета)",
            "Оборуд. текущ. (реестр)", "Оборуд. текущ. (смета)", "Монтаж текущ. (реестр)",
            "Монтаж текущ. (смета)", "Проверка", "Замечания"]
    if PROTO_SHEET in wb.sheetnames:
        pr = wb[PROTO_SHEET]
        if pr.cell(1, 13).value != "Проверка":       # протокол старого образца — пересоздать
            idx = wb.sheetnames.index(PROTO_SHEET)
            wb.remove(pr)
            pr = wb.create_sheet(PROTO_SHEET, idx)
    else:
        pr = wb.create_sheet(PROTO_SHEET)
    if pr.cell(1, 1).value is None:
        for i, h in enumerate(head, 1):
            c = pr.cell(1, i, h)
            c.font, c.fill, c.border = BOLD, HDR_FILL, BOX
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for j, w in enumerate([45, 14, 40, 9, 15, 15, 15, 15, 15, 15, 15, 15, 12, 60], 1):
            pr.column_dimensions[get_column_letter(j)].width = w
        pr.freeze_panes = "A2"
        pr.auto_filter.ref = "A1:N1"
    i = pr.max_row + 1

    def bad(x, y, tol=1.0):
        return y is not None and abs((x or 0) - y) > tol

    for est in estimates:
        poss = [o for k, o in est.items if k == "pos"]
        sb = round(sum(p.base for p in poss), 2)
        sc = round(sum(p.cur for p in poss), 2)
        eq = round(sum(p.cur for p in poss if p.kind == "equip"), 2)
        mo = round(sum(p.cur for p in poss if p.group == "4" and p.kind != "equip"), 2)
        he = est.hdr_equip[1] if est.hdr_equip else None
        hm = est.hdr_mont[1] if est.hdr_mont else None
        t = 10.0 if est.total_from_hdr else 1.0         # шапка дана в тыс. руб.
        flags = [bad(sb, est.total_base, t), bad(sc, est.total_cur, t), bad(eq, he, 10.0), bad(mo, hm, 10.0)]
        # позиции без строки «Всего по позиции» (одиночные материалы) — по их собственной строке
        open_b = sum(p.base for p in poss if not p.closed and (p.parent is None or not p.parent.closed))
        open_c = sum(p.cur for p in poss if not p.closed and (p.parent is None or not p.parent.closed))
        in_doc = (any(flags[:2]) and poss
                  and abs(sb - est.pos_tot[0] - open_b) <= 0.05 and abs(sc - est.pos_tot[1] - open_c) <= 0.05)
        notes = list(est.warnings)
        if in_doc:
            db = round((est.total_base or 0) - sb, 2)
            dc = round((est.total_cur or 0) - sc, 2)
            rub = lambda v: f"{v:,.2f}".replace(",", " ").replace(".", ",")
            notes.insert(0, f"ошибка в самой смете: сумма её строк «Всего по позиции» ({rub(sb)} / {rub(sc)}) "
                            f"не равна её итогу — разница {rub(db)} / {rub(dc)} руб.; в реестр взяты позиции")
        if in_doc and not flags[2]:
            ok = "в смете ≠"
        elif any(flags[:3]):
            ok = "РАСХОЖДЕНИЕ"
        elif flags[3]:
            ok = "монтаж ≠"
        else:
            ok = "ОК"
        LAST_STATS[ok] = LAST_STATS.get(ok, 0) + 1
        row = [est.src, est.key, est.ls_name, len(poss), sb, est.total_base, sc, est.total_cur,
               eq, he, mo, hm, ok, "; ".join(notes)]
        for j, v in enumerate(row, 1):
            c = pr.cell(i, j, v)
            if isinstance(v, float):
                c.number_format = "#,##0.00"
        for j, f in zip((6, 8, 10, 12), flags):
            if f:
                pr.cell(i, j).fill = NEW_FILL
        pr.cell(i, 13).fill = SEC_FILL if ok == "ОК" else (NEW_FILL if ok == "РАСХОЖДЕНИЕ" else WARN_FILL)
        i += 1
    for est in skipped:
        pr.cell(i, 1, est.src)
        pr.cell(i, 2, est.key)
        pr.cell(i, 3, est.ls_name)
        pr.cell(i, 13, "НЕ ДОБАВЛЕНА").fill = NEW_FILL
        pr.cell(i, 14, "ЛС с таким номером уже есть в файле")
        i += 1
    pr.auto_filter.ref = f"A1:N{max(i - 1, 1)}"


def _unwrap_rows(ws, first):
    """Снять «Переносить текст» во всех строках ниже шапки (шапку не трогаем)."""
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            al = c.alignment
            if al is not None and al.wrap_text:
                c.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical, wrap_text=False,
                                        indent=al.indent, text_rotation=al.text_rotation,
                                        shrink_to_fit=al.shrink_to_fit)
    for r in range(first, ws.max_row + 1):               # высота строк — снова автоматическая
        if r in ws.row_dimensions and ws.row_dimensions[r].height is not None:
            ws.row_dimensions[r].height = None


def write_ls(estimates, out_path, upss_code="", object_name="", append_to=None, allow_dupes=False):
    """Создать новый файл или дописать сметы в лист «-ЛС-» существующего файла.
    Возвращает (добавлено, пропущено_дубликатов)."""
    if append_to:
        keep_vba = append_to.lower().endswith(".xlsm")
        wb = openpyxl.load_workbook(append_to, keep_vba=keep_vba)
        if LS_SHEET in wb.sheetnames:
            ws = wb[LS_SHEET]
        else:
            ws = wb.create_sheet(LS_SHEET, 0)
            _init_ls_sheet(ws, upss_code, object_name or estimates[0].object_name)
        if upss_code and not ws["G1"].value:
            ws["G1"] = upss_code
        if object_name and not ws["G2"].value:
            ws["G2"] = object_name
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = LS_SHEET
        _init_ls_sheet(ws, upss_code, object_name or (estimates[0].object_name if estimates else ""))

    ls = LSSheet.detect(ws)
    if append_to:
        _unwrap_rows(ws, ls.hdr + 1)
    if append_to and ls.missing:
        est_warn = "на листе «-ЛС-» нет колонок: " + ", ".join(ls.missing) + " — они не заполняются"
        for e in estimates:
            e.warnings.append(est_warn)
    have = existing_ls_numbers(ls)
    to_add, skipped = [], []
    for e in estimates:
        if e.key in have and not allow_dupes:
            skipped.append(e)
        else:
            to_add.append(e)

    r = _last_data_row(ls) + 1
    first_new = r
    for est in to_add:
        r = _write_estimate(ls, r, est)
    last = max(r - 1, ls.hdr + 1)
    _extend_total_formulas(ls, last)
    _write_protocol(wb, to_add, skipped, mark_new=bool(append_to))
    _stamp(wb)
    wb.save(out_path)
    return to_add, skipped, first_new


# --------------------------------------------------------------------------- запуск
def collect_files(paths, exclude=None):
    files = []
    for p in paths:
        if os.path.isdir(p):
            for root_dir, _dirs, names in os.walk(p):        # без шаблонов: «[ ]» в путях не мешают
                for nm in names:
                    if nm.lower().endswith((".xlsx", ".xlsm", ".pdf")):
                        files.append(os.path.join(root_dir, nm))
        elif os.path.isfile(p):
            files.append(p)
        else:
            files += glob.glob(p)                              # шаблон из командной строки (*.xlsx)
    ex = os.path.abspath(exclude) if exclude else None
    return sorted({f for f in files
                   if not os.path.basename(f).startswith("~$") and os.path.abspath(f) != ex})


def order_key(e):
    """Порядок смет: по номеру ЛС, основная — перед «искл»."""
    k = ls_sort_key(e.ls_no)
    return (k[0], k[1], 1 if e.excl else 0, k[2])


LAST_STATS = {}     # итоги сверки последнего запуска: {"ОК": n, "РАСХОЖДЕНИЕ": n, …}


def run(paths, out, code="", obj="", log=print, append_to=None, allow_dupes=False, backup=True):
    LAST_STATS.clear()
    files = collect_files(paths, exclude=append_to or out)
    if not files:
        log("Файлы смет не найдены.")
        return None
    ests, failed = [], 0
    for f in files:
        try:
            got, errs = parse_file(f, log)
        except Exception as ex:  # noqa
            failed += 1
            log(f"ПРОПУЩЕН  {os.path.basename(f)}: {ex}")
            continue
        if not got and not errs:
            log(f"ПРОПУЩЕН  {os.path.basename(f)}: нет листов со сметой")
        for e in got:
            n = sum(1 for k, _ in e.items if k == "pos")
            log(f"OK   {e.src}  →  ЛС {e.key}, позиций {n}"
                + (f"  [!] {'; '.join(e.warnings)}" if e.warnings else ""))
        for er in errs:
            failed += 1
            log(f"ПРОПУЩЕН ЛИСТ  {os.path.basename(f)} / {er}")
        ests += got
    if not ests:
        return None
    if SORT_BY_LS_NUMBER:
        ests.sort(key=order_key)
    if not code:
        rep_nos = {e.report_no for e in ests if e.report_no}
        if len(rep_nos) == 1:
            code = rep_nos.pop()
            log(f"Шифр УПСС взят из заголовка отчёта: {code}")
    seen = {}
    for e in ests:
        if e.key in seen:
            e.warnings.append(f"номер ЛС совпадает с «{seen[e.key]}»")
        seen.setdefault(e.key, e.src)

    if append_to and backup:
        import shutil
        base, ext = os.path.splitext(append_to)
        bak = f"{base}_резерв{ext}"
        shutil.copy2(append_to, bak)
        log(f"Резервная копия: {bak}")
    try:
        added, skipped, first = write_ls(ests, out, code, obj, append_to=append_to, allow_dupes=allow_dupes)
    except PermissionError:
        log(f"\nНЕ УДАЛОСЬ СОХРАНИТЬ {out}: файл открыт в Excel. Закройте его и повторите.")
        return None
    for e in skipped:
        log(f"НЕ ДОБАВЛЕНА  ЛС {e.key} ({e.src}) — уже есть в файле")
    log(f"\nГотово: {out}")
    log(f"Добавлено смет: {len(added)} (с {first}-й строки листа «-ЛС-»), "
        f"повторов пропущено: {len(skipped)}, ошибок чтения: {failed}. Сверка — на листе «Протокол».")
    return out


SVOD_SHEET = "Свод смет"


def find_ls_columns(ws, max_scan=30):
    """Найти на листе «-ЛС-» строку заголовков и колонки «№ ЛС», базис, текущий уровень.
    -> (строка_заголовка, кол_№ЛС, кол_база, кол_текущ) или исключение."""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        c_ls = c_b = c_c = None
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            t = " ".join(v.lower().replace("ё", "е").split())
            if c_ls is None and re.match(r"^№\s*лс\b", t):
                c_ls = c
            elif c_b is None and "всего" in t and "базисн" in t:
                c_b = c
            elif c_c is None and "всего" in t and "текущ" in t:
                c_c = c
        if c_ls and c_b and c_c:
            return r, c_ls, c_b, c_c
    raise ValueError("на листе «-ЛС-» не найдена строка заголовков с колонками "
                     "«№ ЛС», «Всего затрат в базисном уровне цен», «ВСЕГО затрат в текущем уровне цен»")


def build_svod(path, log=print, backup=True):
    """Лист «Свод смет»: № сметы (= № ЛС, колонка F), Наименование (пока пусто),
    БАЗА = сумма L, ТЕКУЩАЯ = сумма M листа «-ЛС-» по каждому № ЛС."""
    keep_vba = path.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    if LS_SHEET not in wb.sheetnames:
        log(f"В файле нет листа «{LS_SHEET}».")
        return None
    ls = wb[LS_SHEET]
    try:
        hdr, c_ls, c_b, c_c = find_ls_columns(ls)
    except ValueError as ex:
        log(str(ex))
        return None
    L_ls, L_b, L_c = (get_column_letter(x) for x in (c_ls, c_b, c_c))
    log(f"Колонки на «-ЛС-»: № ЛС — {L_ls}, база — {L_b}, текущая — {L_c} (заголовки в строке {hdr})")
    first = hdr + 1
    last = first
    for r in range(ls.max_row, hdr, -1):
        if any(ls.cell(r, c).value not in (None, "") for c in (c_ls, c_b, c_c)):
            last = r
            break

    # одна строка на каждый № ЛС в порядке появления на листе
    order, base, cur, names = [], {}, {}, {}
    re_name = re.compile(r"^\s*Наименование работ документа\s*[-–—]\s*(.*)$", re.I | re.S)
    scan_cols = range(1, min(ls.max_column, 40) + 1)
    for r in range(first, last + 1):
        f = txt(ls.cell(r, c_ls).value)
        if not f:
            continue
        if f not in base:
            order.append(f)
            base[f] = cur[f] = 0.0
        if not names.get(f):
            for c in scan_cols:
                v = ls.cell(r, c).value
                if isinstance(v, str):
                    m = re_name.match(v)
                    if m:
                        nm = re.sub(r"\s*\(смета на исключение\)\s*$", "", " ".join(m.group(1).split()))
                        if nm:
                            names[f] = nm
                        break
        for col, acc in ((c_b, base), (c_c, cur)):
            v = ls.cell(r, col).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                acc[f] += v
    if not order:
        log("На листе «-ЛС-» не заполнена колонка «№ ЛС».")
        return None
    rows = [(f, base[f], cur[f]) for f in order]
    no_name = [f for f in order if not names.get(f)]

    if backup:
        import shutil
        b, ext = os.path.splitext(path)
        shutil.copy2(path, f"{b}_резерв{ext}")

    # пересоздаём лист на том же месте (или сразу после «-ЛС-»)
    if SVOD_SHEET in wb.sheetnames:
        idx = wb.sheetnames.index(SVOD_SHEET)
        wb.remove(wb[SVOD_SHEET])
    else:
        idx = wb.sheetnames.index(LS_SHEET) + 1
    sv = wb.create_sheet(SVOD_SHEET, idx)

    obj = txt(ls["G2"].value)
    sv["A1"] = "Свод смет" + (f" — {obj}" if obj else "")
    sv["A1"].font = Font(bold=True, size=12)
    if txt(ls["G1"].value):
        sv["A2"] = f"Шифр УПСС: {txt(ls['G1'].value)}"
    head = ["№ сметы", "Наименование сметы", "БАЗА", "ТЕКУЩАЯ"]
    notes = [f"№ ЛС («-ЛС-», кол. {L_ls})", "из «Наименование работ документа - …»", f"Всего затрат в базисном уровне цен, руб. (кол. {L_b})",
             f"ВСЕГО затрат в текущем уровне цен, руб. (кол. {L_c})"]
    for j, t in enumerate(head, 1):
        c = sv.cell(4, j, t)
        c.font, c.fill, c.border = BOLD, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        n = sv.cell(3, j, notes[j - 1])
        n.font = Font(italic=True, size=8, color="808080")
        n.alignment = Alignment(horizontal="center", wrap_text=True)
    q = f"'{LS_SHEET}'"
    fr = f"{q}!${L_ls}${first}:${L_ls}${last}"
    r = 5
    for f, b_, c_ in rows:
        cell = sv.cell(r, 1, f)
        cell.border, cell.number_format = BOX, "@"
        c = sv.cell(r, 2, names.get(f))                  # из «Наименование работ документа - …»
        c.border, c.alignment = BOX, Alignment(wrap_text=True, vertical="top")
        for j, col in ((3, L_b), (4, L_c)):
            # точное сравнение текста № ЛС (SUMIF мог бы принять «01-01» за дату)
            c = sv.cell(r, j, f"=SUMPRODUCT(--({fr}=$A{r}),{q}!{col}${first}:{col}${last})")
            c.number_format, c.border = "#,##0.00", BOX
        r += 1
    sv.cell(r, 1, "Итого").font = BOLD
    for j, col in ((3, "C"), (4, "D")):
        c = sv.cell(r, j, f"=SUM({col}5:{col}{r - 1})")
        c.number_format, c.font = "#,##0.00", BOLD
    for j in range(1, 5):
        sv.cell(r, j).border = BOX
        sv.cell(r, j).fill = HDR_FILL
    sv.column_dimensions["A"].width = 22
    sv.column_dimensions["B"].width = 50
    sv.column_dimensions["C"].width = 20
    sv.column_dimensions["D"].width = 20
    sv.freeze_panes = "A5"
    sv.auto_filter.ref = f"A4:D{r - 1}"
    _stamp(wb)
    sv.cell(r + 2, 1, f"Сформировано программой «УПСС · Обработка сметных листов». {COPYRIGHT}").font = \
        Font(italic=True, size=8, color="808080")
    try:
        wb.save(path)
    except PermissionError:
        log(f"НЕ УДАЛОСЬ СОХРАНИТЬ {path}: файл открыт в Excel. Закройте его и повторите.")
        return None
    tb = sum(x[1] for x in rows)
    tc = sum(x[2] for x in rows)
    fmt = lambda v: f"{v:,.2f}".replace(",", " ").replace(".", ",")
    log(f"Свод смет собран: {len(rows)} ЛС. База {fmt(tb)} руб., текущая {fmt(tc)} руб.")
    if no_name:
        log(f"Без наименования (нет строки «Наименование работ документа - …»): {', '.join(no_name)}")
    log(f"Лист «{SVOD_SHEET}» в файле {path}")
    return path


# ── Лист «Итог» (по кнопке «ИТОГ»): суммы по Разделу/Подразделу КР со ссылками на «-ЛС-» ──
ITOG_SHEET = "Итог"
ITOG_ROWS = [                       # (Первый уровень, Второй уровень, Код раздела) — как в шаблоне 4Т
    ('Автоматизированная система коммерческого учета энергоснабжающих ресурсов', 'ИТОГО', '2.2.8'),
    ('Автоматизированная система управления инженерными системами', 'ИТОГО', '2.2.7'),
    ('Благоустройство и озеленение территории', 'ИТОГО', '7.1'),
    ('Внутренняя отделка', 'ИТОГО', '2.1.10'),
    ('Внутренняя отделка', 'полы', '2.1.10.2'),
    ('Внутренняя отделка', 'потолки', '2.1.10.3'),
    ('Внутренняя отделка', 'стены', '2.1.10.1'),
    ('Водоотлив (дренаж)', 'ИТОГО', None),
    ('Временные здания и сооружения', 'ИТОГО', None),
    ('Временные здания и сооружения', 'Конструктивные решения (основание кранов)', '8.2'),
    ('Временные здания и сооружения', 'Организация дорожного движения на период строительства', '8.3'),
    ('Земляные работы', 'ИТОГО', '2.1.1'),
    ('Земляные работы', '', None),
    ('Земляные работы', 'обратная засыпка', '2.1.1.1'),
    ('Крепление котлована', 'ИТОГО', '2.1.2'),
    ('Кровля', 'ИТОГО', '2.1.7'),
    ('Лифтовое оборудование', 'ИТОГО', '2.2.12'),
    ('Мусоропровод', 'ИТОГО', None),
    ('Наружная отделка', 'ИТОГО', '2.1.11'),
    ('Наружные сети и сооружения водоснабжения, водоотведения, теплоснабжения и газоснабжения', 'ИТОГО', '6.1'),
    ('Ненесущие стены и перегородки', 'ИТОГО', '2.1.6'),
    ('Ненесущие стены и перегородки', 'ненесущие стены', '2.1.6.1'),
    ('Ненесущие стены и перегородки', 'перегородки', '2.1.6.2'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'ИТОГО', '2.1.5'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'колонны, пилоны: железобетонные', '2.1.5.1'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'колонны: металлические', '2.1.5.2'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'лестницы', '2.1.5.6'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'перекрытие, покрытие (крыша): железобетонные', '2.1.5.4'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'перекрытие, покрытие (крыша): металлические конструкции (фермы, балки и т.п.)', '2.1.5.5'),
    ('Несущие и самонесущие конструкции надземной части из монолитного железобетона и сборных железобетонных изделий, металлических конструкций', 'стены', '2.1.5.3'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'ИТОГО', '2.1.4'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'колонны (пилоны)', '2.1.4.3'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'лестницы', '2.1.4.6'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'перекрытие, покрытие', '2.1.4.5'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'сваи', '2.1.4.1'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'стены', '2.1.4.4'),
    ('Несущие и самонесущие конструкции подземной части из монолитного железобетона и сборных железобетонных изделий', 'фундамент', '2.1.4.2'),
    ('ОДД на период строительства', 'ИТОГО', '8.3'),
    ('Объекты подсобного и обслуживающего назначения', 'ИТОГО', '3.1'),
    ('Объекты транспортного хозяйства и связи', 'ИТОГО', '5.1'),
    ('Объекты энергетического хозяйства', 'ИТОГО', '4.1'),
    ('Основание под башенный кран', 'ИТОГО', '8.2'),
    ('Подготовка территории строительства', 'ИТОГО', None),
    ('Подготовка территории строительства', 'вырубка деревьев и кустарников', '1.4'),
    ('Подготовка территории строительства', 'компенсационное озеленение', '1.5'),
    ('Подпорные стены', 'ИТОГО', None),
    ('Подпорные стены (гл. 7)', 'ИТОГО', '7.7'),
    ('Проемы:', 'ИТОГО', None),
    ('Проемы:', 'дверные блоки, ворота', '2.1.9'),
    ('Проемы:', 'оконные блоки, витражи, остекление лоджий и балконов', '2.1.8'),
    ('Прочие конструктивные решения', 'ИТОГО', '2.1.12'),
    ('Прочие работы и затраты', 'ИТОГО', '9'),
    ('Прочие работы и затраты', 'затраты на перевозку отходов строительства и сноса, в том числе грунта автотранспортными средствами; затраты на прием отходов строительства и сноса отходополучателями', '9.2'),
    ('Прочие работы и затраты', 'затраты на проведение пусконаладочных работ', '9.1'),
    ('Пусконаладочные работы', 'ИТОГО', None),
    ('Сети связи (внутренние)', 'ИТОГО', '2.2.5'),
    ('Система взимания платы', 'ИТОГО', None),
    ('Система водоснабжения и водоотведения', 'ИТОГО', '2.2.2'),
    ('Система газоанализации', 'ИТОГО', '2.2.10'),
    ('Система газоснабжения', 'ИТОГО', '2.2.9'),
    ('Система отопления', 'ИТОГО', '2.2.3'),
    ('Система отопления', '', '2.2.3.1'),
    ('Система отопления', 'индивидуальный тепловой пункт (ИТП)', '2.2.3.2'),
    ('Система холодоснабжения', 'ИТОГО', '2.2.11'),
    ('Системы безопасности', 'ИТОГО', '2.2.6'),
    ('Системы вентиляции, кондиционирования', 'ИТОГО', '2.2.4'),
    ('Технологическое оборудование', 'ИТОГО', '2.3.1.1'),
    ('Цементация грунтов', 'ИТОГО', None),
    ('Эксплуатация подъемно-транспортных механизмов', 'ИТОГО', '2.1.3'),
    ('Электроснабжение, электроосвещение, молниезащита и заземление', 'ИТОГО', '2.2.1'),
    ('посадка деревьев, кустарников', 'ИТОГО', '7.5'),
    ('ТО 10', None, '2.3.1.4'),
    ('ТО 22', None, '2.3.1.2'),
    ('ТО 0', None, '2.3.1.3'),
]


def _hnorm(v):
    return " ".join(str(v).lower().replace("ё", "е").split()) if isinstance(v, str) else ""


def find_itog_columns(ws, max_scan=30):
    """Колонки «-ЛС-» по заголовкам: Раздел КР, Подраздел КР, № ЛС, база, текущая,
    Доп. сведения по оборудованию (ИО/ТО), НДС, Группа монт., Оборудование."""
    for r in range(1, min(ws.max_row, max_scan) + 1):
        col = {}
        for c in range(1, ws.max_column + 1):
            t = _hnorm(ws.cell(r, c).value)
            if not t:
                continue
            for key, ok in (("a", t == "раздел кр"), ("b", t == "подраздел кр"),
                            ("ls", bool(re.match(r"^№\s*лс\b", t))),
                            ("base", "всего" in t and "базисн" in t), ("cur", "всего" in t and "текущ" in t),
                            ("eqinfo", t.startswith("доп. сведения по оборуд")), ("nds", t == "ндс"),
                            ("grp", t.startswith("группа монт")), ("eq", t == "оборудование")):
                if ok and key not in col:          # первая колонка с таким заголовком (A/B, а не D/E)
                    col[key] = c
        need = ("a", "b", "ls", "base", "cur", "eqinfo", "nds", "grp", "eq")
        if all(k in col for k in need):
            return r, col
        if len(col) >= 5:
            miss = [k for k in need if k not in col]
            raise ValueError("на листе «-ЛС-» нет колонок: " + ", ".join(miss))
    raise ValueError("на листе «-ЛС-» не найдена строка заголовков (Раздел КР, Подраздел КР, № ЛС, …)")


def build_itog(path, log=print, backup=True):
    """Лист «Итог» как в 4Т: текущая/базовая стоимость (всего, строительные, монтажные,
    оборудование, прочие) в тыс.руб. по Разделу/Подразделу КР — формулами SUMIFS на «-ЛС-»."""
    keep_vba = path.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    if LS_SHEET not in wb.sheetnames:
        log(f"В файле нет листа «{LS_SHEET}».")
        return None
    ls = wb[LS_SHEET]
    try:
        hdr, col = find_itog_columns(ls)
    except ValueError as ex:
        log(str(ex))
        return None
    L = {k: get_column_letter(v) for k, v in col.items()}
    log("Колонки на «-ЛС-»: " + ", ".join(f"{n} — {L[k]}" for k, n in (
        ("a", "Раздел КР"), ("b", "Подраздел КР"), ("ls", "№ ЛС"), ("base", "база"), ("cur", "текущая"),
        ("eqinfo", "ИО/ТО"), ("nds", "НДС"), ("grp", "группа"), ("eq", "оборудование"))))
    first = hdr + 1
    last = max(first, ls.max_row)

    # проверка на Python: что есть на «-ЛС-» и что не попадёт в итог
    seen, no_kr = [], [0.0, 0.0]
    known = {(a, b or "") for a, b, _ in ITOG_ROWS}
    known_a = {a for a, b, _ in ITOG_ROWS if b == "ИТОГО"}
    for r in range(first, last + 1):
        vb, vc = ls.cell(r, col["base"]).value, ls.cell(r, col["cur"]).value
        num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
        if not num(vb) and not num(vc):
            continue
        a, b = txt(ls.cell(r, col["a"]).value), txt(ls.cell(r, col["b"]).value)
        if not a:
            no_kr[0] += num(vb)
            no_kr[1] += num(vc)
            continue
        for key in ((a, "ИТОГО"), (a, b)):
            if key[1] and key not in known and key not in seen:
                seen.append(key)
    extra = seen                                     # разделы/подразделы, которых нет в шаблоне

    if backup:
        import shutil
        bn, ext = os.path.splitext(path)
        shutil.copy2(path, f"{bn}_резерв{ext}")

    if ITOG_SHEET in wb.sheetnames:
        idx = wb.sheetnames.index(ITOG_SHEET)
        wb.remove(wb[ITOG_SHEET])
    else:
        idx = wb.sheetnames.index(SVOD_SHEET) + 1 if SVOD_SHEET in wb.sheetnames \
            else wb.sheetnames.index(LS_SHEET) + 1
    it = wb.create_sheet(ITOG_SHEET, idx)

    head = ["Первый уровень", "Второй уровень",
            "Текущая стоимость по ЛС тыс.руб.", "Текущая стоимость (строительные) тыс.руб.",
            "Текущая стоимость (монтажные) тыс.руб.", "Текущая стоимость (оборудование) тыс.руб.",
            "Текущая стоимость (Прочие) тыс.руб.",
            "Базовая стоимость по ЛС тыс.руб.", "Базовая стоимость (строительные) тыс.руб.",
            "Базовая стоимость (монтажные) тыс.руб.", "Базовая стоимость (оборудование) тыс.руб.",
            "Базовая стоимость (Прочие) тыс.руб.", "Код раздела", "№ показателя"]
    for j, t in enumerate(head, 1):
        c = it.cell(1, j, t)
        c.font, c.fill, c.border = BOLD, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    it.row_dimensions[1].height = 60
    upss = txt(ls["G1"].value)

    q = f"'{LS_SHEET}'"
    rng = lambda k: f"{q}!${L[k]}${first}:${L[k]}${last}"
    R_A, R_B, R_LS, R_EI, R_NDS, R_G, R_EQ = (rng(k) for k in ("a", "b", "ls", "eqinfo", "nds", "grp", "eq"))
    sums = {"cur": rng("cur"), "base": rng("base")}

    def f_sum(what, crit):
        return f"=ROUND(SUMIFS({sums[what]},{crit})/1000,3)"

    body = [(a, b, code) for a, b, code in ITOG_ROWS if not a.startswith("ТО ")]
    body += [(a, b, "") for a, b in extra]
    to_rows = [(a, b, code) for a, b, code in ITOG_ROWS if a.startswith("ТО ")]
    r = 3
    first_body = r
    body_rows = {(a, b): first_body + i for i, (a, b, _) in enumerate(body)}
    total_rows = {a: body_rows[(a, "ИТОГО")] for a, b, _ in body if b == "ИТОГО"}
    mismatch_rows = []
    for a, b, code in body:
        it.cell(r, 1, a)
        it.cell(r, 2, b)
        child_rows = [row for (child_a, child_b), row in body_rows.items()
                      if child_a == a and child_b not in ("", "ИТОГО")]
        reconciliation_rows = [row for (child_a, child_b), row in body_rows.items()
                               if child_a == a and child_b != "ИТОГО"]
        if not b and a in total_rows:
            for c0 in (3, 8):
                for c in range(c0, c0 + 5):
                    col_letter = get_column_letter(c)
                    children = ",".join(f"{col_letter}{child}" for child in child_rows)
                    it.cell(r, c, f"=ROUND({col_letter}{total_rows[a]}-SUM({children}),3)")
        else:
            crit = f"{R_A},$A{r}" + ("" if b == "ИТОГО" else f",{R_B},$B{r}")
            for c0, what in ((3, "cur"), (8, "base")):
                it.cell(r, c0, f_sum(what, crit))
                it.cell(r, c0 + 1, f"=ROUND({get_column_letter(c0)}{r}-{get_column_letter(c0 + 2)}{r}-"
                                   f"{get_column_letter(c0 + 3)}{r}-{get_column_letter(c0 + 4)}{r},3)")
                it.cell(r, c0 + 2, f_sum(what, f'{crit},{R_G},"4",{R_EQ},""'))      # монтажные (группа 4, не оборуд.)
                it.cell(r, c0 + 3, f_sum(what, f'{crit},{R_EQ},"*О*"'))              # оборудование
                it.cell(r, c0 + 4, f_sum(what, f'{crit},{R_LS},"09-*"'))             # прочие (глава 9)
        it.cell(r, 13, code or None)
        it.cell(r, 14, upss or None)
        if b == "ИТОГО":
            for j in (1, 2):
                it.cell(r, j).font = BOLD
            if reconciliation_rows:
                comparisons = []
                for c in range(3, 13):
                    col_letter = get_column_letter(c)
                    child_refs = ",".join(f"${col_letter}${child}" for child in reconciliation_rows)
                    comparisons.append(
                        f"OR(ROUND(${col_letter}{r},3)<>ROUND(SUM({child_refs}),3),"
                        f"ROUND(${col_letter}{r},0)<>ROUND(SUM({child_refs}),0))"
                    )
                mismatch_rows.append((r, f"=OR({','.join(comparisons)})"))
        if (a, b) in extra:
            for j in range(1, 15):
                it.cell(r, j).fill = EXCL_FILL
        r += 1
    for row, formula in mismatch_rows:
        it.conditional_formatting.add(
            f"A{row}:N{row}",
            FormulaRule(formula=[formula], fill=ITOG_MISMATCH_FILL, font=ITOG_MISMATCH_FONT),
        )
    last_body = r - 1
    for a, b, code in to_rows:                       # ТО по ставке НДС
        nds = a.split()[-1]
        it.cell(r, 1, a)
        for j in range(3, 13):
            it.cell(r, j, 0)
        for j, what in ((6, "cur"), (11, "base")):
            it.cell(r, j, f'=ROUND(SUMIFS({sums[what]},{R_EI},"ТО",{R_NDS},"{"" if nds == "0" else nds}")/1000,3)')
        it.cell(r, 13, code)
        it.cell(r, 14, upss or None)
        r += 1
    last_row = r - 1

    it.cell(2, 1, "ОБЩИЙ ИТОГ")
    for j in range(3, 13):
        L_ = get_column_letter(j)
        if j in (4, 9):
            it.cell(2, j, f"=ROUND({get_column_letter(j - 1)}2-{get_column_letter(j + 1)}2-"
                          f"{get_column_letter(j + 2)}2-{get_column_letter(j + 3)}2,3)")
        else:
            it.cell(2, j, f'=SUMIF(B{first_body}:B{last_body},"ИТОГО",{L_}{first_body}:{L_}{last_body})')
    it.cell(2, 14, upss or None)
    for j in range(1, 15):
        it.cell(2, j).font = BOLD
        it.cell(2, j).fill = LS_FILL
    for rr in range(2, last_row + 1):
        for j in range(1, 15):
            c = it.cell(rr, j)
            c.border = BOX
            if 3 <= j <= 12:
                c.number_format = "#,##0.000"
        it.cell(rr, 1).alignment = Alignment(wrap_text=True, vertical="top")
        it.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")
    for k, w in zip("ABCDEFGHIJKLMN", (66, 25, 18, 18, 18, 18, 18, 18, 18, 18, 18, 18, 12, 15)):
        it.column_dimensions[k].width = w
    it.freeze_panes = "C3"
    it.auto_filter.ref = f"A1:N{last_row}"
    it.cell(last_row + 2, 1, f"Сформировано программой «УПСС · Обработка сметных листов». {COPYRIGHT}").font = \
        Font(italic=True, size=8, color="808080")
    _stamp(wb)
    try:
        wb.save(path)
    except PermissionError:
        log(f"НЕ УДАЛОСЬ СОХРАНИТЬ {path}: файл открыт в Excel. Закройте его и повторите.")
        return None
    fmt = lambda v: f"{v:,.2f}".replace(",", " ").replace(".", ",")
    log(f"ИТОГ собран: {len(body)} строк разделов + {len(to_rows)} строки ТО.")
    if extra:
        log("Нет в шаблоне — добавлены в конец (выделены цветом): " + "; ".join(f"{a} | {b}" for a, b in extra))
    if no_kr[0] or no_kr[1]:
        log(f"ВНИМАНИЕ: строки без «Раздел КР» не попали в ИТОГ — база {fmt(no_kr[0])} руб., "
            f"текущая {fmt(no_kr[1])} руб. Заполните колонку {L['a']} на «-ЛС-».")
    log(f"Лист «{ITOG_SHEET}» в файле {path}")
    return path


# ── Грунт (кнопка «Грунт»): колонки T…AC листа «-ЛС-» + лист «Грунт» ──────────────
GRUNT_SHEET = "Грунт"
G_FILL_BLUE = PatternFill("solid", start_color="FF0000FF", end_color="FF0000FF")
G_FILL_YEL = PatternFill("solid", start_color="FFFFFF99", end_color="FFFFFF99")
G_FILL_HEAD = PatternFill("solid", start_color="FF00B0F0", end_color="FF00B0F0")
# std-колонка «-ЛС-» → правило (префиксы шифра); ×100 — расценки на 100 м3
G_DIG = {
    "U": ("3.1-1-", "3.1-2-", "3.1-3-"),                                   # Разработка в отвал
    "V": ("3.1-9-", "3.1-10-", "3.1-11-", "3.1-12-", "3.1-13-"),           # без погрузки
    "X": ("3.1-5-", "3.1-6-", "3.1-7-", "3.1-24-"),                        # с погрузкой
    "W": ("3.1-49-", "3.1-50-", "3.1-51-", "3.1-52-", "3.1-55-"),          # вручную
    "Y": ("3.1-14-", "3.1-15-", "3.1-16-", "3.1-53-"),                     # Засыпка
}
G_MATERIAL = ("1.1-1-766", "1.1-1-3709", "1.1-1-4474", "1.1-1-1511", "1.1-1-160")   # точный шифр
# лист «Грунт»: колонка → std-колонка «-ЛС-»
G_DST = [("Утилизация грунта", "AC"), ("Перевозка грунта", "T"), ("Разработка в отвал", "U"),
         ("Разработка без погрузки", "V"), ("Разработка вручную", "W"), ("Разработка с погрузкой", "X"),
         ("Засыпка", "Y"), ("Мусор", "AA"), ("Перевозка мусора", "AB")]


def _style_grunt_headers(ws, cols, hdr):
    """cols: {фактическая колонка: std-буква}. Заголовок — синий, «МАТЕРИАЛ» — жёлтый."""
    for c, k in cols.items():
        h = ws.cell(hdr, c, GRUNT[k])
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = BOX
        if k == "Z":
            h.fill, h.font = G_FILL_YEL, Font(color="3535FF", bold=True)
        else:
            h.fill, h.font = G_FILL_BLUE, Font(color="FFFFFF", bold=True)
        if ws.column_dimensions[get_column_letter(c)].width in (None, 13.0):
            ws.column_dimensions[get_column_letter(c)].width = 14


def _grunt_kind(code, name, qty):
    """-> {std-буква: значение} для строки «-ЛС-» (правила кнопки «Грунт»)."""
    out = {}
    if not code:
        return out
    code = str(code).strip()
    nl = (name or "").lower()
    if "добавляется на" in nl or "добавлять на" in nl:
        return out
    num = isinstance(qty, (int, float)) and not isinstance(qty, bool)
    x100 = qty * 100 if num else qty
    if "перевозка" in nl and "грунт" in nl:
        out["T"] = qty
    for k, pref in G_DIG.items():
        if code.startswith(pref):
            out[k] = x100
            return out
    if code.startswith("15.1-") and "грунт" not in nl:
        out["AA"] = qty
    elif code.startswith("15.2-") and "мусор" in nl:
        out["AB"] = qty
    elif code.startswith("15.1-") and "грунт" in nl:
        out["AC"] = qty
    elif code in G_MATERIAL:
        out["Z"] = qty
    return out


def build_grunt(path, log=print, backup=True):
    """Кнопка «Грунт»: заполнить колонки грунта на «-ЛС-» (T…AC) и собрать лист «Грунт» со ссылками."""
    keep_vba = path.lower().endswith(".xlsm")
    wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    if LS_SHEET not in wb.sheetnames:
        log(f"В файле нет листа «{LS_SHEET}».")
        return None
    ws = wb[LS_SHEET]
    ls = LSSheet.detect(ws)
    need = {6: "№ ЛС", 7: "№ п/п", 8: "Шифр", 9: "Наименование", 11: "Кол-во"}
    miss = [t for c, t in need.items() if not ls.col(c)]
    if miss:
        log("На листе «-ЛС-» нет колонок: " + ", ".join(miss))
        return None
    hdr, first = ls.hdr, ls.hdr + 1
    # колонки грунта: найденные по заголовку, иначе — стандартные T…AC (или справа от таблицы)
    gcol = {}
    free = max(ws.max_column, 18) + 2
    for k in GRUNT:
        std = openpyxl.utils.column_index_from_string(k)
        c = ls.cmap.get(std) if ls.cmap.get(std) else None
        if not c:
            c = std if all(ws.cell(r, std).value in (None, "") for r in range(1, hdr + 1)) and \
                std not in ls.cmap.values() else free
            if c == free:
                free += 1
        gcol[k] = c
        for r in range(1, hdr):                                  # заголовки «над шапкой» (старая версия) — убрать
            if _norm(ws.cell(r, c).value or "") == _norm(GRUNT[k]):
                ws.cell(r, c).value = None
    last = _last_data_row(ls)

    if backup:
        import shutil
        bn, ext = os.path.splitext(path)
        shutil.copy2(path, f"{bn}_резерв{ext}")

    # 1) правая часть «-ЛС-»
    c_code, c_name, c_qty = ls.col(8), ls.col(9), ls.col(11)
    c_ls, c_no, c_a, c_g = ls.col(6), ls.col(7), ls.col(1), ls.col(7)
    hits = []                        # (строка, № ЛС, № п/п, раздел)
    stat = {k: 0 for k in GRUNT}
    sect = ""
    for r in range(first, last + 1):
        for k in GRUNT:
            ws.cell(r, gcol[k]).value = None
        g = ws.cell(r, c_g).value
        if isinstance(g, str) and RE_SECTION.match(g) and g.lower().startswith("раздел"):
            sect = RE_SECTION.match(g).group(1).strip()
            continue
        name = ws.cell(r, c_name).value
        res = _grunt_kind(ws.cell(r, c_code).value, str(name) if name is not None else "", ws.cell(r, c_qty).value)
        if not res:
            continue
        for k, v in res.items():
            cell = ws.cell(r, gcol[k], v)
            cell.number_format = "#,##0.00"
            stat[k] += 1
        if any(k != "Z" for k in res):
            kr = txt(ws.cell(r, c_a).value) if c_a else ""
            hits.append((r, ws.cell(r, c_ls).value, ws.cell(r, c_no).value, kr or sect))
    _style_grunt_headers(ws, {c: k for k, c in gcol.items()}, hdr)
    for k, c in gcol.items():
        L_ = get_column_letter(c)
        s = ws.cell(hdr - 1 if hdr > 1 else 1, c)
        if k == "Z":
            s.value = None
            continue
        s.value = f"=SUBTOTAL(9,{L_}{first}:{L_}{max(last, first)})"
        s.fill, s.font, s.number_format = G_FILL_YEL, Font(bold=True), "#,##0.00"

    # 2) лист «Грунт»
    if GRUNT_SHEET in wb.sheetnames:
        idx = wb.sheetnames.index(GRUNT_SHEET)
        wb.remove(wb[GRUNT_SHEET])
    else:
        idx = len(wb.sheetnames)
    gs = wb.create_sheet(GRUNT_SHEET, idx)
    head = ["№ сметы", "номер по порядку", "раздел"] + [t for t, _ in G_DST]
    for j, t in enumerate(head, 1):
        c = gs.cell(1, j, t)
        c.font, c.fill, c.border = Font(bold=True), G_FILL_HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    q = f"'{LS_SHEET}'"
    fmt0 = "#,##0.00;-#,##0.00;"
    dr = 4
    for r, f, no, sec in hits:
        gs.cell(dr, 1, f).border = BOX
        gs.cell(dr, 2, no).border = BOX
        gs.cell(dr, 3, sec or None).border = BOX
        for j, (_, k) in enumerate(G_DST, 4):
            c = gs.cell(dr, j, f"={q}!{get_column_letter(gcol[k])}{r}")      # нули скрыты форматом
            c.number_format, c.border = fmt0, BOX
        dr += 1
    last_g = max(dr - 1, 4)
    for j in range(1, len(head) + 1):
        c = gs.cell(2, j)
        if j >= 4:
            L_ = get_column_letter(j)
            c.value = f"=SUBTOTAL(9,{L_}4:{L_}{last_g})"
            c.number_format = fmt0
        c.font, c.fill, c.border = Font(bold=True), G_FILL_YEL, BOX
        c.alignment = Alignment(horizontal="center")
    gs.auto_filter.ref = f"A3:{get_column_letter(len(head))}{last_g}"
    gs.freeze_panes = "A4"
    gs.column_dimensions["A"].width = 14
    gs.column_dimensions["B"].width = 10
    gs.column_dimensions["C"].width = 40
    for j in range(4, len(head) + 1):
        gs.column_dimensions[get_column_letter(j)].width = 13.3
    gs.cell(last_g + 2, 1, f"Сформировано программой «УПСС · Обработка сметных листов». {COPYRIGHT}").font = \
        Font(italic=True, size=8, color="808080")
    _stamp(wb)
    try:
        wb.save(path)
    except PermissionError:
        log(f"НЕ УДАЛОСЬ СОХРАНИТЬ {path}: файл открыт в Excel. Закройте его и повторите.")
        return None
    log("ГРУНТ: колонки на «-ЛС-» — " + ", ".join(f"{GRUNT[k]} {get_column_letter(c)}" for k, c in gcol.items()))
    log("ГРУНТ: заполнено строк — " + ", ".join(f"{GRUNT[k]}: {n}" for k, n in stat.items() if n) if any(stat.values())
        else "ГРУНТ: работ по грунту/мусору в сметах не найдено.")
    log(f"Грунт собран: {len(hits)} строк на листе «{GRUNT_SHEET}». Файл {path}")
    return path


def open_file(path, log=print):
    """Открыть файл в Excel. Пути из диалогов Windows приходят с «/» (в т.ч. сетевые //сервер/…) —
    приводим к виду «\\\\сервер\\…», иначе Windows их не находит."""
    p = os.path.normpath(os.path.abspath(path)) if not path.startswith(("//", "\\\\")) else \
        os.path.normpath(path).replace("/", "\\")
    try:
        if hasattr(os, "startfile"):
            os.startfile(p)
    except OSError as ex:
        log(f"Файл сохранён, но открыть его автоматически не удалось ({ex}). Откройте вручную: {p}")


def gui():
    import json
    import queue
    import threading
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    # ---------- палитра и шрифты ----------
    C = {
        "bg": "#F3F5F9", "card": "#FFFFFF", "line": "#E2E7EF", "text": "#1E2532", "muted": "#6B7484",
        "accent": "#2457D6", "accent_h": "#1C47B3", "accent_soft": "#E8EEFC",
        "head": "#14213D", "head_sub": "#9FB0D0",
        "ok": "#1E9E5A", "warn": "#C98A10", "bad": "#D2433B", "info": "#2457D6",
        "log_bg": "#101828", "log_fg": "#D5DCE8", "log_dim": "#7C8AA5",
        "chip_ok": "#E6F6EE", "chip_warn": "#FDF3DF", "chip_bad": "#FCE9E8", "chip_info": "#E8EEFC",
    }
    root = tk.Tk()
    root.title("УПСС — обработка сметных листов   ·   © KAS-11")
    root.geometry("1180x760")
    root.minsize(1000, 640)
    root.configure(bg=C["bg"])

    import tkinter.font as tkfont
    fams = set(tkfont.families())
    base = next((f for f in ("Segoe UI", "Inter", "Arial", "DejaVu Sans") if f in fams), "TkDefaultFont")
    mono = next((f for f in ("Cascadia Mono", "Consolas", "DejaVu Sans Mono") if f in fams), "TkFixedFont")
    F = {
        "h1": (base, 17, "bold"), "h1s": (base, 10), "h2": (base, 11, "bold"), "b": (base, 10),
        "s": (base, 9), "btn": (base, 10, "bold"), "big": (base, 12, "bold"), "num": (base, 16, "bold"),
        "mono": (mono, 9),
    }

    st = ttk.Style(root)
    try:
        st.theme_use("clam")
    except tk.TclError:
        pass
    st.configure("Accent.Horizontal.TProgressbar", troughcolor=C["line"], background=C["accent"],
                 bordercolor=C["line"], lightcolor=C["accent"], darkcolor=C["accent"], thickness=6)
    st.configure("TCheckbutton", background=C["card"], foreground=C["text"], font=F["b"])
    st.map("TCheckbutton", background=[("active", C["card"])])
    st.configure("TEntry", fieldbackground="#FAFBFD", bordercolor=C["line"], lightcolor=C["line"], padding=6)
    st.configure("Vertical.TScrollbar", gripcount=0, background="#D5DBE5", troughcolor="#F4F6FA",
                 bordercolor="#F4F6FA", lightcolor="#D5DBE5", darkcolor="#D5DBE5", arrowcolor="#8B93A3",
                 arrowsize=11)
    st.map("Vertical.TScrollbar", background=[("active", "#B9C2D0")])
    st.configure("Log.Vertical.TScrollbar", gripcount=0, background="#2A3550", troughcolor=C["log_bg"],
                 bordercolor=C["log_bg"], lightcolor="#2A3550", darkcolor="#2A3550", arrowcolor="#7C8AA5",
                 arrowsize=11)

    # ---------- вход по паролю ----------
    def password_dialog(mode="login"):
        """mode: login — ввод пароля; create — задать новый; change — сменить. Возвращает True при успехе."""
        dlg = tk.Toplevel(root)
        dlg.title("УПСС — вход" if mode == "login" else "УПСС — пароль")
        dlg.configure(bg=C["head"])
        dlg.resizable(False, False)
        res = {"ok": False, "tries": 0}
        box = tk.Frame(dlg, bg=C["card"], padx=28, pady=24)
        box.pack(padx=1, pady=(0, 1), fill="both")
        top = tk.Frame(dlg, bg=C["head"], padx=28, pady=16)
        top.pack(before=box, fill="x")
        tk.Label(top, text="УПСС  ·  Обработка сметных листов", font=F["h2"], fg="white", bg=C["head"]).pack(anchor="w")
        tk.Label(top, text=COPYRIGHT, font=F["s"], fg=C["head_sub"], bg=C["head"]).pack(anchor="w", pady=(2, 0))
        titles = {"login": "Введите пароль", "create": "Задайте пароль для запуска программы",
                  "change": "Смена пароля"}
        tk.Label(box, text=titles[mode], font=F["h2"], fg=C["text"], bg=C["card"]).pack(anchor="w")
        fields = []
        labels = {"login": ["Пароль"], "create": ["Новый пароль", "Повторите пароль"],
                  "change": ["Текущий пароль", "Новый пароль", "Повторите новый пароль"]}[mode]
        for lab in labels:
            tk.Label(box, text=lab, font=F["s"], fg=C["muted"], bg=C["card"]).pack(anchor="w", pady=(12, 2))
            e = ttk.Entry(box, show="•", width=34, font=F["b"])
            e.pack(fill="x")
            fields.append(e)
        msg = tk.Label(box, text="", font=F["s"], fg=C["bad"], bg=C["card"], wraplength=300, justify="left")
        msg.pack(anchor="w", pady=(8, 0))

        def ok(_e=None):
            vals = [f.get() for f in fields]
            if mode == "login":
                if check_password(vals[0]):
                    res["ok"] = True
                    dlg.destroy()
                    return
                res["tries"] += 1
                if res["tries"] >= 3:
                    dlg.destroy()
                    return
                msg.configure(text=f"Неверный пароль. Осталось попыток: {3 - res['tries']}")
                fields[0].delete(0, "end")
                return
            if mode == "change" and not check_password(vals[0]):
                msg.configure(text="Текущий пароль неверный")
                return
            new, rep_ = vals[-2], vals[-1]
            if len(new) < 4:
                msg.configure(text="Пароль должен быть не короче 4 символов")
                return
            if new != rep_:
                msg.configure(text="Пароли не совпадают")
                return
            try:
                store_password(new)
            except Exception as ex:  # noqa
                msg.configure(text=f"Не удалось сохранить пароль: {ex}")
                return
            res["ok"] = True
            dlg.destroy()

        row = tk.Frame(box, bg=C["card"])
        row.pack(fill="x", pady=(14, 0))
        go = tk.Label(row, text="Войти" if mode == "login" else "Сохранить", font=F["btn"], bg=C["accent"],
                      fg="white", padx=22, pady=8, cursor="hand2")
        go.pack(side="right")
        go.bind("<Button-1>", ok)
        go.bind("<Enter>", lambda e: go.configure(bg=C["accent_h"]))
        go.bind("<Leave>", lambda e: go.configure(bg=C["accent"]))
        cancel = tk.Label(row, text="Отмена", font=F["btn"], bg=C["card"], fg=C["muted"], padx=14, pady=8,
                          cursor="hand2")
        cancel.pack(side="right", padx=6)
        cancel.bind("<Button-1>", lambda e: dlg.destroy())
        dlg.bind("<Return>", ok)
        dlg.bind("<Escape>", lambda e: dlg.destroy())
        dlg.update_idletasks()
        w, h = dlg.winfo_reqwidth(), dlg.winfo_reqheight()
        dlg.geometry(f"+{(dlg.winfo_screenwidth() - w) // 2}+{(dlg.winfo_screenheight() - h) // 3}")
        fields[0].focus_set()
        dlg.grab_set()
        root.wait_window(dlg)
        return res["ok"]

    root.withdraw()
    if not PASSWORD_HASH:
        if not password_dialog("create"):
            root.destroy()
            return
    elif not password_dialog("login"):
        root.destroy()
        return
    root.deiconify()

    # ---------- настройки (запоминаем между запусками) ----------
    cfg_path = os.path.join(os.path.expanduser("~"), ".upss_ls.json")
    try:
        cfg = json.load(open(cfg_path, encoding="utf-8"))
    except Exception:  # noqa
        cfg = {}

    def save_cfg():
        try:
            json.dump(cfg, open(cfg_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        except Exception:  # noqa
            pass

    state = {"files": [], "target": cfg.get("target"), "result": None, "busy": False}
    q = queue.Queue()

    # ---------- виджеты-помощники ----------
    def card(parent, title, step=None):
        outer = tk.Frame(parent, bg=C["line"])
        inner = tk.Frame(outer, bg=C["card"], padx=16, pady=14)
        inner.pack(fill="both", expand=True, padx=1, pady=1)
        hdr = tk.Frame(inner, bg=C["card"])
        hdr.pack(fill="x", pady=(0, 10))
        if step:
            tk.Label(hdr, text=str(step), font=F["s"] + ("bold",), fg="white", bg=C["accent"],
                     width=2, padx=2).pack(side="left", padx=(0, 8))
        tk.Label(hdr, text=title, font=F["h2"], fg=C["text"], bg=C["card"]).pack(side="left")
        return outer, inner, hdr

    def button(parent, text, cmd, kind="secondary", width=None):
        pal = {
            "primary": (C["accent"], "white", C["accent_h"]),
            "secondary": (C["accent_soft"], C["accent"], "#D9E3FA"),
            "ghost": (C["card"], C["muted"], "#F0F2F6"),
        }[kind]
        b = tk.Label(parent, text=text, font=F["big"] if kind == "primary" else F["btn"], bg=pal[0], fg=pal[1],
                     padx=18 if kind == "primary" else 12, pady=10 if kind == "primary" else 6, cursor="hand2")
        if width:
            b.configure(width=width)
        b._pal = pal
        b._enabled = True
        b.bind("<Enter>", lambda e: b._enabled and b.configure(bg=pal[2]))
        b.bind("<Leave>", lambda e: b.configure(bg=pal[0] if b._enabled else "#E9ECF1"))
        b.bind("<Button-1>", lambda e: b._enabled and cmd())

        def enable(on=True):
            b._enabled = on
            b.configure(bg=pal[0] if on else "#E9ECF1", fg=pal[1] if on else "#A3AAB6",
                        cursor="hand2" if on else "arrow")
        b.enable = enable
        return b

    # ---------- шапка ----------
    head = tk.Frame(root, bg=C["head"], padx=24, pady=16)
    head.pack(fill="x")
    tk.Label(head, text="УПСС  ·  Обработка сметных листов", font=F["h1"], fg="white", bg=C["head"]).pack(anchor="w")
    tk.Label(head, text="Сметы Excel и PDF  →  лист «-ЛС-» по форме 4Т  ·  Свод смет  ·  сверка с итогами смет",
             font=F["h1s"], fg=C["head_sub"], bg=C["head"]).pack(anchor="w", pady=(2, 0))

    foot = tk.Frame(root, bg=C["bg"], padx=22, pady=6)
    foot.pack(side="bottom", fill="x")
    tk.Label(foot, text=COPYRIGHT, font=F["s"], fg=C["muted"], bg=C["bg"]).pack(side="left")
    chg = tk.Label(foot, text="Сменить пароль", font=F["s"], fg=C["accent"], bg=C["bg"], cursor="hand2")
    chg.pack(side="right")

    def change_pw(_e=None):
        if password_dialog("change"):
            messagebox.showinfo("Пароль", "Пароль изменён.")
    chg.bind("<Button-1>", change_pw)

    body = tk.Frame(root, bg=C["bg"], padx=20, pady=16)
    body.pack(fill="both", expand=True)
    body.columnconfigure(0, weight=5, uniform="c")
    body.columnconfigure(1, weight=6, uniform="c")
    body.rowconfigure(0, weight=1)

    left = tk.Frame(body, bg=C["bg"])
    left.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
    right = tk.Frame(body, bg=C["bg"])
    right.grid(row=0, column=1, sticky="nsew", padx=(10, 0))

    # ---------- 1. Сметы ----------
    c1o, c1, c1h = card(left, "Сметы", 1)
    c1o.pack(fill="both", expand=True)
    cnt_lbl = tk.Label(c1h, text="не выбраны", font=F["s"], fg=C["muted"], bg=C["card"])
    cnt_lbl.pack(side="right")
    bar1 = tk.Frame(c1, bg=C["card"])
    bar1.pack(fill="x")
    lst_wrap = tk.Frame(c1, bg=C["line"])
    lst_wrap.pack(fill="both", expand=True, pady=(10, 0))
    lst = tk.Listbox(lst_wrap, font=F["b"], bd=0, highlightthickness=0, activestyle="none",
                     selectbackground=C["accent_soft"], selectforeground=C["text"], fg=C["text"],
                     bg="#FAFBFD", selectmode="extended")
    sb1 = ttk.Scrollbar(lst_wrap, orient="vertical", command=lst.yview)
    lst.configure(yscrollcommand=sb1.set)
    sb1.pack(side="right", fill="y", padx=(0, 1), pady=1)
    lst.pack(fill="both", expand=True, padx=1, pady=1)
    hint = tk.Label(c1, text="Excel (.xlsx, .xlsm) и PDF. В одном файле может быть много ЛС — на листах или подряд.",
                    font=F["s"], fg=C["muted"], bg=C["card"], wraplength=420, justify="left")
    hint.pack(anchor="w", pady=(8, 0))

    def refresh_files():
        lst.delete(0, "end")
        n_files = 0
        for p in state["files"]:
            if os.path.isdir(p):
                k = len(collect_files([p]))
                n_files += k
                name = os.path.basename(p.rstrip("/").rstrip("\\")) or p
                lst.insert("end", f"  ПАПКА   {name}   ·  {k} файл(ов)")
            else:
                n_files += 1
                ext = os.path.splitext(p)[1].lower()
                tag = "PDF  " if ext == ".pdf" else "EXCEL"
                lst.insert("end", f"  {tag}   {os.path.basename(p)}")
        cnt_lbl.configure(text=f"{n_files} файл(ов)" if n_files else "не выбраны",
                          fg=C["accent"] if n_files else C["muted"])
        update_buttons()

    def add_files():
        sel = filedialog.askopenfilenames(
            title="Сметы", initialdir=cfg.get("dir_in") or None,
            filetypes=[("Сметы Excel/PDF", "*.xlsx *.xlsm *.pdf"), ("Excel", "*.xlsx *.xlsm"), ("PDF", "*.pdf")])
        if sel:
            cfg["dir_in"] = os.path.dirname(sel[0])
            for s in sel:
                if s not in state["files"]:
                    state["files"].append(s)
            refresh_files()

    def add_dir():
        d = filedialog.askdirectory(title="Папка со сметами", initialdir=cfg.get("dir_in") or None)
        if d:
            cfg["dir_in"] = d
            if d not in state["files"]:
                state["files"].append(d)
            refresh_files()

    def remove_sel():
        for i in reversed(lst.curselection()):
            del state["files"][i]
        refresh_files()

    def clear_files():
        state["files"].clear()
        refresh_files()

    button(bar1, "＋ Файлы", add_files).pack(side="left")
    button(bar1, "＋ Папка", add_dir).pack(side="left", padx=6)
    button(bar1, "Очистить", clear_files, "ghost").pack(side="right")
    button(bar1, "Убрать", remove_sel, "ghost").pack(side="right", padx=4)
    lst.bind("<Delete>", lambda e: remove_sel())

    # ---------- 2. Куда записать ----------
    c2o, c2, _ = card(left, "Куда записать", 2)
    c2o.pack(fill="x", pady=(14, 0))
    mode = tk.StringVar(value=cfg.get("mode", "new"))
    seg = tk.Frame(c2, bg=C["line"])
    seg.pack(fill="x")
    seg_btns = {}

    def set_mode(m):
        mode.set(m)
        cfg["mode"] = m
        for k, b in seg_btns.items():
            on = k == m
            b.configure(bg=C["card"] if on else C["line"], fg=C["accent"] if on else C["muted"],
                        font=F["btn"] if on else F["b"])
        tgt_row.pack_forget()
        if m == "append":
            tgt_row.pack(fill="x", pady=(10, 0))
        update_buttons()

    for k, t in (("new", "Новый файл"), ("append", "Добавить к существующему")):
        b = tk.Label(seg, text=t, pady=7, cursor="hand2", font=F["b"])
        b.pack(side="left", fill="x", expand=True, padx=1, pady=1)
        b.bind("<Button-1>", lambda e, k=k: set_mode(k))
        seg_btns[k] = b

    tgt_row = tk.Frame(c2, bg=C["card"])
    tgt_lbl = tk.Label(tgt_row, text="файл не выбран", font=F["s"], fg=C["muted"], bg=C["card"],
                       anchor="w", justify="left", wraplength=330)
    tgt_lbl.pack(side="left", fill="x", expand=True)

    def show_target():
        t = state["target"]
        if not t:
            tgt_lbl.configure(text="файл не выбран", fg=C["muted"])
            return
        info = ""
        try:
            wb = openpyxl.load_workbook(t, read_only=True)
            if LS_SHEET in wb.sheetnames:
                nums = existing_ls_numbers(openpyxl.load_workbook(t)[LS_SHEET])
                info = f"   ·  уже ЛС: {len(nums)}"
            else:
                info = "   ·  листа «-ЛС-» нет, будет создан"
        except Exception as ex:  # noqa
            info = f"   ·  не открывается: {ex}"
        tgt_lbl.configure(text=os.path.basename(t) + info, fg=C["text"])

    def pick_target():
        t = filedialog.askopenfilename(title="Файл, к которому добавить сметы",
                                       initialdir=cfg.get("dir_out") or None,
                                       filetypes=[("Excel", "*.xlsx *.xlsm")])
        if t:
            state["target"] = t
            cfg["target"] = t
            cfg["dir_out"] = os.path.dirname(t)
            show_target()
            update_buttons()

    button(tgt_row, "Выбрать…", pick_target).pack(side="right")
    dupes = tk.BooleanVar(value=False)
    ttk.Checkbutton(c2, text="Добавлять ЛС, даже если такой номер уже есть", variable=dupes).pack(anchor="w", pady=(10, 0))

    # ---------- 3. Параметры ----------
    c3o, c3, _ = card(left, "Параметры объекта", 3)
    c3o.pack(fill="x", pady=(14, 0))
    grid3 = tk.Frame(c3, bg=C["card"])
    grid3.pack(fill="x")
    grid3.columnconfigure(1, weight=1)
    tk.Label(grid3, text="Шифр УПСС", font=F["s"], fg=C["muted"], bg=C["card"]).grid(row=0, column=0, sticky="w")
    e_code = ttk.Entry(grid3, font=F["b"], width=14)
    e_code.grid(row=1, column=0, sticky="w", padx=(0, 12))
    tk.Label(grid3, text="Наименование объекта", font=F["s"], fg=C["muted"], bg=C["card"]).grid(row=0, column=1, sticky="w")
    e_obj = ttk.Entry(grid3, font=F["b"])
    e_obj.grid(row=1, column=1, sticky="ew")
    tk.Label(c3, text="Можно не заполнять — шифр и объект возьмутся из смет.", font=F["s"], fg=C["muted"],
             bg=C["card"]).pack(anchor="w", pady=(6, 0))

    # ---------- правая колонка: действия, итоги, журнал ----------
    act = tk.Frame(right, bg=C["bg"])
    act.pack(fill="x")
    b_start = button(act, "▶  Начать обработку", lambda: start(), "primary")
    b_start.pack(side="left")
    b_svod = button(act, "Σ  Собрать Свод смет", lambda: svod(), "secondary")
    b_svod.pack(side="left", padx=(10, 0), ipady=4)
    b_itog = button(act, "≡  ИТОГ", lambda: itog(), "secondary")
    b_itog.pack(side="left", padx=(10, 0), ipady=4)
    b_grunt = button(act, "⛏  Грунт", lambda: grunt(), "secondary")
    b_grunt.pack(side="left", padx=10, ipady=4)
    b_open = button(act, "Открыть результат", lambda: open_file(state["result"], say), "ghost")
    b_open.pack(side="right", ipady=4)

    prog = ttk.Progressbar(right, style="Accent.Horizontal.TProgressbar", mode="determinate", maximum=100, value=0)
    prog.pack(fill="x", pady=(14, 0))
    status = tk.Label(right, text="Готово к работе", font=F["s"], fg=C["muted"], bg=C["bg"], anchor="w")
    status.pack(fill="x", pady=(4, 0))

    chips = tk.Frame(right, bg=C["bg"])
    chips.pack(fill="x", pady=(10, 0))
    chip_w = {}
    for key, title, col, bgc in (("all", "ЛС обработано", C["info"], C["chip_info"]),
                                 ("ОК", "Сошлись", C["ok"], C["chip_ok"]),
                                 ("warn", "Монтаж / в смете ≠", C["warn"], C["chip_warn"]),
                                 ("РАСХОЖДЕНИЕ", "Расхождения", C["bad"], C["chip_bad"])):
        fr = tk.Frame(chips, bg=bgc, padx=14, pady=8)
        fr.pack(side="left", fill="x", expand=True, padx=(0, 8) if key != "РАСХОЖДЕНИЕ" else 0)
        n = tk.Label(fr, text="—", font=F["num"], fg=col, bg=bgc)
        n.pack(anchor="w")
        tk.Label(fr, text=title, font=F["s"], fg=C["muted"], bg=bgc).pack(anchor="w")
        chip_w[key] = n

    def set_chips(stats):
        if not stats:
            for n in chip_w.values():
                n.configure(text="—")
            return
        chip_w["all"].configure(text=str(sum(stats.values())))
        chip_w["ОК"].configure(text=str(stats.get("ОК", 0)))
        chip_w["warn"].configure(text=str(stats.get("монтаж ≠", 0) + stats.get("в смете ≠", 0)))
        chip_w["РАСХОЖДЕНИЕ"].configure(text=str(stats.get("РАСХОЖДЕНИЕ", 0)))

    logo, logi, logh = card(right, "Журнал")
    logo.pack(fill="both", expand=True, pady=(14, 0))
    logi.configure(bg=C["log_bg"])
    logh.configure(bg=C["log_bg"])
    for w in logh.winfo_children():
        w.configure(bg=C["log_bg"], fg=C["log_dim"])
    log_txt = tk.Text(logi, font=F["mono"], bg=C["log_bg"], fg=C["log_fg"], bd=0, highlightthickness=0,
                      wrap="word", insertbackground=C["log_fg"], padx=4, pady=2)
    sb2 = ttk.Scrollbar(logi, orient="vertical", command=log_txt.yview, style="Log.Vertical.TScrollbar")
    log_txt.configure(yscrollcommand=sb2.set)
    sb2.pack(side="right", fill="y")
    log_txt.pack(fill="both", expand=True)
    log_txt.tag_configure("ok", foreground="#5FD39A")
    log_txt.tag_configure("warn", foreground="#F2C14E")
    log_txt.tag_configure("bad", foreground="#FF7B72")
    log_txt.tag_configure("dim", foreground=C["log_dim"])
    log_txt.tag_configure("head", foreground="white", font=(mono, 9, "bold"))

    def write_log(s):
        tag = None
        if s.startswith("OK"):
            tag = "ok" if "[!]" not in s else "warn"
        elif s.startswith(("ПРОПУЩЕН", "НЕ ", "Traceback")) or "ошибк" in s.lower() and "ошибок чтения: 0" not in s:
            tag = "bad"
        elif s.startswith(("   …", "Резервная", "Колонки")):
            tag = "dim"
        elif s.startswith(("Готово", "Свод смет собран", "ИТОГ собран", "Грунт собран", "Добавлено смет")):
            tag = "head"
        log_txt.insert("end", s + "\n", tag)
        log_txt.see("end")

    def say(s):            # вызывается из рабочего потока
        q.put(("log", s))

    def poll():
        try:
            while True:
                kind, val = q.get_nowait()
                if kind == "log":
                    for line in str(val).split("\n"):
                        write_log(line)
                        if line.startswith("OK"):
                            status.configure(text=line[5:90])
                elif kind == "done":
                    finish(*val)
        except queue.Empty:
            pass
        root.after(80, poll)

    # ---------- логика ----------
    def update_buttons():
        busy = state["busy"]
        ok_start = bool(state["files"]) and (mode.get() == "new" or bool(state["target"]))
        b_start.enable(ok_start and not busy)
        b_svod.enable(not busy)
        b_itog.enable(not busy)
        b_grunt.enable(not busy)
        b_open.enable(bool(state["result"]) and not busy)

    def set_busy(on, text=""):
        state["busy"] = on
        if on:
            prog.configure(mode="indeterminate", value=0)
            prog.start(12)
            status.configure(text=text, fg=C["accent"])
        else:
            prog.stop()
            prog.configure(mode="determinate", value=100)
        update_buttons()

    def worker(fn, *args):
        try:
            res = fn(*args)
            q.put(("done", (res, None)))
        except Exception:  # noqa
            q.put(("done", (None, traceback.format_exc())))

    def finish(res, err):
        kind = state.get("job")
        set_busy(False)
        if err:
            write_log(err)
            status.configure(text="Ошибка — подробности в журнале", fg=C["bad"])
            return
        if not res:
            status.configure(text="Не выполнено — см. журнал", fg=C["bad"])
            return
        state["result"] = res
        if kind == "run":
            set_chips(dict(LAST_STATS))
            bad = LAST_STATS.get("РАСХОЖДЕНИЕ", 0)
            status.configure(text=("Готово. Есть расхождения — см. лист «Протокол»" if bad
                                   else "Готово. Все сметы сошлись с итогами"),
                             fg=C["bad"] if bad else C["ok"])
        elif kind == "itog":
            status.configure(text="Лист «Итог» собран", fg=C["ok"])
        elif kind == "grunt":
            status.configure(text="Грунт собран", fg=C["ok"])
        else:
            status.configure(text="Свод смет собран", fg=C["ok"])
        update_buttons()
        if messagebox.askyesno("Готово", "Открыть файл с результатом?"):
            open_file(res, say)

    def start():
        if state["busy"] or not state["files"]:
            return
        if mode.get() == "append":
            if not state["target"]:
                pick_target()
                if not state["target"]:
                    return
            out, app = state["target"], state["target"]
        else:
            out = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="-ЛС-.xlsx",
                                               initialdir=cfg.get("dir_out") or None,
                                               filetypes=[("Excel", "*.xlsx")])
            if not out:
                return
            cfg["dir_out"] = os.path.dirname(out)
            app = None
            state["target"] = out
            cfg["target"] = out
        cfg["code"], cfg["obj"] = e_code.get().strip(), e_obj.get().strip()
        save_cfg()
        log_txt.delete("1.0", "end")
        set_chips(None)
        state["job"] = "run"
        set_busy(True, "Обработка… большие PDF читаются 1–3 минуты")
        threading.Thread(target=worker, daemon=True,
                         args=(run, list(state["files"]), out, e_code.get().strip(), e_obj.get().strip(), say,
                               app, dupes.get())).start()

    def svod():
        if state["busy"]:
            return
        t = state["target"] if state["target"] and os.path.exists(state["target"]) else None
        if not t or not messagebox.askyesno("Свод смет", f"Собрать свод в файле\n{os.path.basename(t)}?\n\n"
                                                         "«Нет» — выбрать другой файл."):
            t = filedialog.askopenfilename(title="Файл с листом «-ЛС-» для свода",
                                           initialdir=cfg.get("dir_out") or None,
                                           filetypes=[("Excel", "*.xlsx *.xlsm")])
            if not t:
                return
        log_txt.delete("1.0", "end")
        state["job"] = "svod"
        set_busy(True, "Сбор свода смет…")
        threading.Thread(target=worker, args=(build_svod, t, say), daemon=True).start()

    def itog():
        if state["busy"]:
            return
        t = state["target"] if state["target"] and os.path.exists(state["target"]) else None
        if not t or not messagebox.askyesno("ИТОГ", f"Собрать лист «Итог» в файле\n{os.path.basename(t)}?\n\n"
                                                    "«Нет» — выбрать другой файл."):
            t = filedialog.askopenfilename(title="Файл с листом «-ЛС-» для ИТОГа",
                                           initialdir=cfg.get("dir_out") or None,
                                           filetypes=[("Excel", "*.xlsx *.xlsm")])
            if not t:
                return
        log_txt.delete("1.0", "end")
        state["job"] = "itog"
        set_busy(True, "Сбор листа «Итог»…")
        threading.Thread(target=worker, args=(build_itog, t, say), daemon=True).start()

    def grunt():
        if state["busy"]:
            return
        t = state["target"] if state["target"] and os.path.exists(state["target"]) else None
        if not t or not messagebox.askyesno("Грунт", f"Заполнить грунт в файле\n{os.path.basename(t)}?\n\n"
                                                     "«Нет» — выбрать другой файл."):
            t = filedialog.askopenfilename(title="Файл с листом «-ЛС-» для грунта",
                                           initialdir=cfg.get("dir_out") or None,
                                           filetypes=[("Excel", "*.xlsx *.xlsm")])
            if not t:
                return
        log_txt.delete("1.0", "end")
        state["job"] = "grunt"
        set_busy(True, "Грунт: заполнение колонок и листа…")
        threading.Thread(target=worker, args=(build_grunt, t, say), daemon=True).start()

    # стартовое состояние
    e_code.insert(0, cfg.get("code", ""))
    e_obj.insert(0, cfg.get("obj", ""))
    set_mode(mode.get())
    show_target()
    refresh_files()
    write_log("Добавьте сметы и нажмите «Начать обработку».")
    write_log("Сверка каждой ЛС с её итогами — на листе «Протокол» результата.")

    def on_close():
        save_cfg()
        root.destroy()
    root.protocol("WM_DELETE_WINDOW", on_close)
    poll()
    root.mainloop()


def cli_login():
    import getpass
    if not PASSWORD_HASH:
        print("Пароль ещё не задан. Запустите программу без параметров (окно) и задайте пароль.")
        sys.exit(1)
    for _ in range(3):
        if check_password(getpass.getpass("Пароль: ")):
            return
        print("Неверный пароль.")
    sys.exit(1)


def main():
    if len(sys.argv) == 1:
        gui()
        return
    print("УПСС · Обработка сметных листов.  " + COPYRIGHT)
    cli_login()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="*", help="файлы смет .xlsx и/или папки")
    ap.add_argument("--svod", help="собрать лист «Свод смет» из листа «-ЛС-» указанного файла")
    ap.add_argument("--itog", help="собрать лист «Итог» из листа «-ЛС-» указанного файла")
    ap.add_argument("--grunt", help="заполнить грунт на «-ЛС-» и собрать лист «Грунт» в указанном файле")
    ap.add_argument("-o", "--out", default="-ЛС-.xlsx", help="итоговый файл (новый)")
    ap.add_argument("-a", "--append", help="добавить сметы в существующий файл (лист «-ЛС-»)")
    ap.add_argument("--dupes", action="store_true", help="добавлять ЛС, даже если номер уже есть")
    ap.add_argument("--code", default="", help="шифр УПСС, напр. 5-5-1-1")
    ap.add_argument("--object", default="", help="наименование объекта")
    a = ap.parse_args()
    if (a.svod or a.itog or a.grunt) and not a.paths:
        if a.grunt:
            build_grunt(a.grunt)
        if a.svod:
            build_svod(a.svod)
        if a.itog:
            build_itog(a.itog)
        return
    if not a.paths:
        ap.error("укажите файлы смет или --svod файл.xlsx")
    if a.append:
        run(a.paths, a.append, a.code, a.object, append_to=a.append, allow_dupes=a.dupes)
    else:
        run(a.paths, a.out, a.code, a.object)


if __name__ == "__main__":
    main()
